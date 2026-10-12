"""Synthetic regression fixtures; no matter files or mail are used."""
from __future__ import annotations

import importlib.util
import sys
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def script(name):
    sys.path.insert(0, str(ROOT / "scripts"))
    try:
        spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / (name + ".py"))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.pop(0)


class MetadataIntegrityTests(unittest.TestCase):
    def test_bundled_template_retains_valid_compatibility_namespaces(self):
        with zipfile.ZipFile(ROOT / "assets" / "neutral-letterhead.docx") as package:
            for name in package.namelist():
                if name.endswith(".xml"):
                    for element in etree.fromstring(package.read(name)).iter():
                        for prefix in element.get("{" + MC + "}Ignorable", "").split():
                            self.assertIn(prefix, element.nsmap, name)
    def test_scrub_preserves_namespace_bindings_for_ignorable_prefixes(self):
        helper = script("scrub_docx_metadata")
        xml = (f'<w:document xmlns:w="{W}" xmlns:mc="{MC}" '
               'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
               'mc:Ignorable="w14"><w:body><w:p w:rsidR="00112233"/></w:body></w:document>')
        root = etree.fromstring(helper.scrub_xml("word/document.xml", xml.encode()))
        for prefix in root.get("{" + MC + "}Ignorable").split():
            self.assertIn(prefix, root.nsmap)
        self.assertNotIn("{" + W + "}rsidR", root.find(".//{" + W + "}p").attrib)

    def test_generated_template_retains_valid_compatibility_namespaces(self):
        builder = script("build_neutral_template")
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "neutral.docx"
            builder.build(output)
            Document(output)
            with zipfile.ZipFile(output) as package:
                for name in package.namelist():
                    if name.endswith(".xml"):
                        root = etree.fromstring(package.read(name))
                        for element in root.iter():
                            for prefix in element.get("{" + MC + "}Ignorable", "").split():
                                self.assertIn(prefix, element.nsmap, name)

    def test_scrub_invalid_xml_leaves_no_partial_output(self):
        helper = script("scrub_docx_metadata")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source, output = folder / "fictional.docx", folder / "scrubbed.docx"
            Document().save(source)
            with zipfile.ZipFile(source) as package:
                parts = {name: package.read(name) for name in package.namelist()}
            parts["docProps/core.xml"] = b"<broken"
            with zipfile.ZipFile(source, "w") as package:
                for name, data in parts.items():
                    package.writestr(name, data)
            with self.assertRaises((ValueError, etree.XMLSyntaxError, ParseError)):
                helper.scrub(source, output)
            self.assertFalse(output.exists())

    def test_custom_properties_are_removed_with_package_references(self):
        helper = script("scrub_docx_metadata")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source, output = folder / "fictional.docx", folder / "scrubbed.docx"
            Document().save(source)
            with zipfile.ZipFile(source) as package:
                parts = {name: package.read(name) for name in package.namelist()}
            root = etree.fromstring(parts["_rels/.rels"])
            etree.SubElement(root, "{" + helper.REL_NS + "}Relationship", Id="fictionalCustom", Target="docProps/custom.xml", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties")
            parts["_rels/.rels"] = etree.tostring(root)
            root = etree.fromstring(parts["[Content_Types].xml"])
            etree.SubElement(root, "{" + helper.CONTENT_NS + "}Override", PartName="/docProps/custom.xml", ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml")
            parts["[Content_Types].xml"] = etree.tostring(root)
            parts["docProps/custom.xml"] = b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"/>'
            with zipfile.ZipFile(source, "w") as package:
                for name, data in parts.items():
                    package.writestr(name, data)
            helper.scrub(source, output)
            with zipfile.ZipFile(output) as package:
                self.assertNotIn("docProps/custom.xml", package.namelist())
                self.assertNotIn(b"docProps/custom.xml", package.read("_rels/.rels"))
                self.assertNotIn(b"docProps/custom.xml", package.read("[Content_Types].xml"))


def fictional_pdf(path, page_texts):
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in page_texts:
        page_id, stream_id = len(objects) + 1, len(objects) + 2
        kids.append(f"{page_id} 0 R")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {stream_id} 0 R >>".encode())
        stream = f"BT /F1 11 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects[1] = f"<< /Type /Pages /Count {len(kids)} /Kids [{' '.join(kids)}] >>".encode()
    data, offsets = b"%PDF-1.4\n", [0]
    for number, item in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{number} 0 obj\n".encode() + item + b"\nendobj\n"
    xref = len(data)
    data += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    data += f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(data)


@unittest.skipUnless(shutil.which("pdftotext") and shutil.which("pdfinfo"), "Poppler commands unavailable")
class PdfCompletenessTests(unittest.TestCase):
    def test_core_extraction_also_flags_blank_final_page(self):
        from lawyeredos.document_helpers import extract_pdf
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "fictional.pdf"
            fictional_pdf(source, ["Fictional complete page. " * 8, ""])
            result = extract_pdf(source, folder / "core.txt")
            self.assertTrue(result["needs_ocr"])
            self.assertEqual(result["sparse_pages"], [2])
    def test_dense_first_page_does_not_mask_blank_final_page(self):
        from lawyeredos.document_helpers import extract_pdf
        helper = script("extract_pdf")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "fictional.pdf"
            fictional_pdf(source, ["Fictional complete page. " * 8, ""])
            args = type("Args", (), {"input": str(source), "output": str(folder / "script.txt"), "pdftotext": "pdftotext", "pdfinfo": "pdfinfo", "minimum_characters": 80})()
            self.assertEqual(helper.run(args)["status"], "review")
            result = extract_pdf(source, folder / "core.txt")
            self.assertTrue(result["needs_ocr"])
            self.assertEqual(result["sparse_pages"], [2])

    def test_review_cache_preserves_blank_final_page_coverage(self):
        from lawyeredos.content_review import _pdf_units
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "fictional.pdf"
            fictional_pdf(source, ["Fictional complete page. " * 8, ""])
            for attempt in range(2):
                with self.subTest(attempt=attempt):
                    units, coverage = _pdf_units(source, folder / "scratch")
                    self.assertEqual(coverage["pages"], 2)
                    self.assertEqual(coverage["empty_pages"], [2])
                    self.assertFalse(coverage["complete"])

    def test_fully_textual_pages_pass_completeness(self):
        from lawyeredos.document_helpers import extract_pdf
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "fictional.pdf"
            fictional_pdf(source, ["Fictional first page. " * 8, "Fictional second page. " * 8])
            result = extract_pdf(source, folder / "text.txt")
            self.assertFalse(result["needs_ocr"])


class PdfAssessmentTests(unittest.TestCase):
    def test_missing_final_page_requires_review(self):
        from lawyeredos.pdf_assessment import assess_pdf_text
        result = assess_pdf_text("Fictional page content. " * 8 + "\f", 2)
        self.assertEqual(result["missing_pages"], [2])
        self.assertFalse(result["page_count_verified"])
        self.assertTrue(result["needs_ocr"])

    def test_unknown_page_count_requires_review(self):
        from lawyeredos.pdf_assessment import assess_pdf_text
        self.assertTrue(assess_pdf_text("Fictional content. " * 8 + "\f", None)["needs_ocr"])

    def test_blank_and_sparse_pages_remain_counted(self):
        from lawyeredos.pdf_assessment import assess_pdf_text
        result = assess_pdf_text("Fictional content. " * 8 + "\fShort\f\f", 3)
        self.assertEqual(result["extracted_pages"], 3)
        self.assertEqual(result["sparse_pages"], [2, 3])
        self.assertTrue(result["needs_ocr"])


if __name__ == "__main__":
    unittest.main()

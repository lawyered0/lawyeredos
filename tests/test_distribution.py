from __future__ import annotations

import importlib.util
import hashlib
import io
import json
import re
import shutil
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr
from pathlib import Path
from urllib.parse import unquote

from lawyeredos.codex_setup import _projectize_skill


ROOT = Path(__file__).resolve().parents[1]


def load_script(name: str):
    path = ROOT / "scripts" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(f"practice_{name}", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


class DistributionTests(unittest.TestCase):
    def test_codex_projectization_rewrites_nested_skill_references(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "lawyeredos-work-on-matter"
            _projectize_skill(
                ROOT / "skills" / "practice-work-on-matter",
                destination,
                "lawyeredos-work-on-matter",
            )
            nested = (destination / "references" / "mail-intake.md").read_text(encoding="utf-8")
            boundaries = (
                destination / "references" / "capability-boundaries.md"
            ).read_text(encoding="utf-8")
            self.assertIn('".lawyeredos/lawyeredos-run" -m lawyeredos.mail_intake plan', nested)
            self.assertIn(".lawyeredos/current-codex/scripts/extract_pdf.py", boundaries)
            self.assertNotIn("LAWYEREDOS_ROOT", nested + boundaries)
            self.assertNotIn("LAWYEREDOS_PYTHON", nested + boundaries)
            self.assertNotIn("LAWYEREDOS_CONFIG", nested + boundaries)

    def test_setup_entrypoints_reject_old_python_before_core_import(self):
        for name in ("setup_wizard", "setup_codex"):
            text = (ROOT / "scripts" / f"{name}.py").read_text(encoding="utf-8")
            self.assertLess(
                text.index("require_supported_python()"),
                text.index("from lawyeredos"),
            )
            module = load_script(name)
            error = io.StringIO()
            with redirect_stderr(error), self.assertRaises(SystemExit) as raised:
                module.require_supported_python((3, 9))
            self.assertEqual(raised.exception.code, 2)
            self.assertIn("requires Python 3.11 or newer", error.getvalue())

    def test_manifest_and_six_skills(self):
        manifest_path = ROOT / ".claude-plugin" / "plugin.json"
        self.assertTrue(manifest_path.is_file())
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertEqual(manifest["name"], "lawyeredos")
        self.assertFalse(manifest["name"].startswith(("claude-", "anthropic-", "cc-plugin-")))
        self.assertEqual(manifest["agents"], ["./agents/legal-review-lens.md"])
        skills = sorted((ROOT / "skills").glob("*/SKILL.md"))
        self.assertEqual(len(skills), 6)

    def test_plugin_cross_references_are_namespaced_and_self_contained(self):
        sibling_names = {
            "practice-work-on-matter",
            "practice-matter-tracker",
            "practice-letterhead",
            "practice-contract-redline",
            "practice-calendar-sync",
            "practice-adversarial-panel",
        }
        combined_body = ""
        for path in (ROOT / "skills").glob("*/SKILL.md"):
            text = path.read_text(encoding="utf-8")
            body = text.split("---", 2)[-1]
            combined_body += body
            for name in sibling_names:
                self.assertIsNone(
                    re.search(rf"(?<!lawyeredos:)\b{re.escape(name)}\b", body),
                    f"unqualified sibling reference {name} in {path.relative_to(ROOT)}",
                )
        self.assertIn("lawyeredos:practice-work-on-matter", combined_body)
        self.assertIn("lawyeredos:practice-matter-tracker", combined_body)
        self.assertIn("lawyeredos:practice-adversarial-panel", combined_body)
        self.assertIn("lawyeredos:legal-review-lens", combined_body)
        self.assertNotIn("panel-lens", combined_body)

    def test_public_skills_preserve_portable_practice_guardrails(self):
        work = (ROOT / "skills" / "practice-work-on-matter" / "SKILL.md").read_text(encoding="utf-8")
        tracker = (ROOT / "skills" / "practice-matter-tracker" / "SKILL.md").read_text(encoding="utf-8")
        panel = (ROOT / "skills" / "practice-adversarial-panel" / "SKILL.md").read_text(encoding="utf-8")
        letter = (ROOT / "skills" / "practice-letterhead" / "SKILL.md").read_text(encoding="utf-8")
        redline = (ROOT / "skills" / "practice-contract-redline" / "SKILL.md").read_text(encoding="utf-8")
        self.assertIn("A closed matter remains read-only", work)
        self.assertIn("without a preliminary full orientation", work)
        self.assertIn("brief.truncated: true", work)
        self.assertIn("From address alone does not prove", work)
        self.assertIn("No located hit in the local register fields", tracker)
        self.assertIn("Disposition every material finding", panel)
        self.assertIn("panel concern alone never", panel)
        self.assertIn("Correct every FAIL in original prose", letter)
        self.assertIn("Hash the exact source before generation", redline)
        self.assertRegex(letter, r"never overwrite newer\s+work")
        self.assertRegex(redline, r"never overwrite newer\s+work")

    def test_bootstrap_setup_skill_is_project_local(self):
        setup_skill = ROOT / ".claude" / "skills" / "lawyeredos-setup" / "SKILL.md"
        text = setup_skill.read_text(encoding="utf-8")
        self.assertIn("name: lawyeredos-setup", text)
        self.assertNotIn("disable-model-invocation: true", text)
        self.assertIn("ordinary request", text)
        self.assertIn("scripts/setup_wizard.py", text)
        self.assertIn("--apply --yes --json", text)
        self.assertIn("explicit approval", text)
        self.assertIn("Do not edit the target", text)

    def test_codex_bootstrap_is_project_local_and_separate(self):
        setup_skill = ROOT / ".agents" / "skills" / "lawyeredos-setup" / "SKILL.md"
        text = setup_skill.read_text(encoding="utf-8")
        self.assertIn("name: lawyeredos-setup", text)
        self.assertIn("scripts/setup_codex.py", text)
        self.assertIn("~/.codex", text)
        self.assertIn("--require-new-workspace", text)
        self.assertTrue((ROOT / "CODEX.md").is_file())

    def test_lawyeredos_branding_and_namespace(self):
        manifest = json.loads((ROOT / ".claude-plugin" / "plugin.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["name"], "lawyeredos")
        self.assertEqual(manifest["displayName"], "LawyeredOS")
        project = (ROOT / "pyproject.toml").read_text(encoding="utf-8")
        self.assertIn('name = "lawyeredos"', project)
        self.assertTrue((ROOT / "src" / "lawyeredos" / "__init__.py").is_file())

        legacy = re.compile("|".join([
            "claude" + r"[_-]" + "legal" + r"[_-]" + "practice",
            "legal" + "-practice:",
            "CLAUDE" + "_LEGAL_PRACTICE_CONFIG",
            "LEGAL" + r"_PRACTICE_(?:GMAIL|CALENDAR)_TOKEN",
        ]), re.IGNORECASE)
        findings = []
        for path in ROOT.rglob("*"):
            if any(part in {".venv", "qa", "vendor", "__pycache__", ".pytest_cache"} for part in path.parts):
                continue
            if not path.is_file() or path.suffix.lower() not in {".md", ".txt", ".toml", ".json", ".py", ".yaml", ".yml"}:
                continue
            if legacy.search(path.read_text(encoding="utf-8", errors="replace")):
                findings.append(str(path.relative_to(ROOT)))
        self.assertEqual(findings, [])

    def test_host_support_is_explicit(self):
        host_support = (ROOT / "docs" / "HOST-SUPPORT.md").read_text(encoding="utf-8")
        self.assertIn("Claude Desktop 2.31226.0", host_support)
        self.assertIn("required tested target", host_support)
        self.assertIn("Desktop Code", host_support)
        self.assertIn("`claude --desktop --plugin-dir ...`", host_support)
        self.assertNotIn("Cowork", host_support)

    def test_no_absolute_home_paths_or_common_tokens(self):
        patterns = [
            re.compile(r"/(?:Users|home)/[^/\s]+/"),
            re.compile(r"\bsk-[A-Za-z0-9_-]{20,}\b"),
            re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
        ]
        for path in ROOT.rglob("*"):
            if any(part in {".venv", "qa", "__pycache__", ".pytest_cache"} for part in path.parts):
                continue
            if not path.is_file() or path.suffix.lower() not in {".md", ".txt", ".toml", ".json", ".py", ".yaml", ".yml"}:
                continue
            text = path.read_text(encoding="utf-8", errors="replace")
            for pattern in patterns:
                self.assertIsNone(pattern.search(text), f"{pattern.pattern} in {path.relative_to(ROOT)}")

    def test_template_metadata_and_content(self):
        template = ROOT / "assets" / "neutral-letterhead.docx"
        if not template.exists():
            self.skipTest("neutral template has not been generated")
        with zipfile.ZipFile(template) as package:
            self.assertNotIn("docProps/custom.xml", package.namelist())
            self.assertFalse(any(name.startswith("customXml/") for name in package.namelist()))
            combined = b"\n".join(package.read(name) for name in package.namelist() if name.endswith(".xml"))
        text = combined.decode("utf-8", errors="replace")
        self.assertNotRegex(text, r"/(?:Users|home)/")
        self.assertIn("example.invalid", text)
        self.assertNotIn("rsid", text)
        self.assertNotIn("docId", text)
        self.assertNotIn("Microsoft Macintosh Word", text)
        self.assertNotIn("lastModifiedBy", text)
        self.assertNotIn("<dc:creator", text)
        self.assertNotIn("<dcterms:created", text)
        self.assertNotIn("<dcterms:modified", text)

    def test_upstream_license_retained(self):
        license_path = ROOT / "vendor" / "legal-redline-tools" / "LICENSE"
        text = license_path.read_text(encoding="utf-8")
        self.assertIn("MIT License", text)
        self.assertIn("Chris Sheehan", text)

    def test_gpl_license_and_separate_third_party_notice(self):
        license_text = (ROOT / "LICENSE").read_text(encoding="utf-8")
        self.assertIn("GNU GENERAL PUBLIC LICENSE", license_text)
        self.assertIn("Version 3, 29 June 2007", license_text)
        project = (ROOT / "pyproject.toml").read_text(encoding="utf-8")
        self.assertIn('license = "GPL-3.0-or-later"', project)
        copyright_text = (ROOT / "COPYRIGHT").read_text(encoding="utf-8")
        self.assertIn("Copyright (C) 2026 Lawyered", copyright_text)
        decision = (ROOT / "LICENSE-DECISION.md").read_text(encoding="utf-8")
        self.assertIn("Copyright (C) 2026", decision)
        self.assertIn("AI assistance", decision)
        notices = (ROOT / "THIRD_PARTY_NOTICES.md").read_text(encoding="utf-8")
        self.assertIn("legal-redline-tools 0.2.0", notices)
        self.assertIn("MIT License", notices)
        self.assertIn("eef61a0cb6afe2867fde70335c179d4a6ad298c8", notices)
        modifications = (
            ROOT / "vendor" / "legal-redline-tools" / "VENDOR-MODIFICATIONS.md"
        ).read_text(encoding="utf-8")
        self.assertIn("eef61a0cb6afe2867fde70335c179d4a6ad298c8", modifications)
        provenance = (ROOT / "PROVENANCE.md").read_text(encoding="utf-8")
        self.assertIn("created by the owner with AI assistance", provenance)
        self.assertIn("not a legal conclusion", provenance)

    def test_accept_tracked_changes_end_to_end(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")
        vendor = ROOT / "vendor" / "legal-redline-tools"
        sys.path.insert(0, str(vendor))
        try:
            from legal_redline import apply_redlines
        finally:
            sys.path.pop(0)
        accepter = load_script("accept_tracked_changes")
        with tempfile.TemporaryDirectory() as temporary:
            temporary = Path(temporary)
            original = temporary / "original.docx"
            redlined = temporary / "redlined.docx"
            accepted = temporary / "accepted.docx"
            document = Document()
            document.add_paragraph("The delivery period is thirty days.")
            document.save(original)
            apply_redlines(
                str(original), str(redlined),
                [{"type": "replace", "old": "thirty days", "new": "forty-five days"}],
            )
            accepter.accept(redlined, accepted)
            with zipfile.ZipFile(accepted) as package:
                xml = package.read("word/document.xml").decode("utf-8")
            self.assertNotIn("<w:ins", xml)
            self.assertNotIn("<w:del", xml)
            self.assertIn("forty-five days", Document(accepted).paragraphs[0].text)

    def test_letter_applies_configured_branding_label(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")
        creator = load_script("create_letter")
        with tempfile.TemporaryDirectory() as temporary:
            temporary = Path(temporary)
            template = temporary / "neutral-letterhead.docx"
            shutil.copy2(ROOT / "assets" / "neutral-letterhead.docx", template)
            manifest = temporary / "template-manifest.json"
            manifest.write_text(json.dumps({
                "schema_version": 1, "algorithm": "sha256", "template": template.name,
                "sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
            }), encoding="utf-8")
            body = temporary / "body.txt"
            output = temporary / "letter.docx"
            body.write_text("This is a fictional test letter.\n", encoding="utf-8")
            args = type("Args", (), {
                "template": str(template),
                "template_manifest": str(manifest),
                "body_file": str(body),
                "output": str(output),
                "organization": "Northstar Legal Lab",
                "date": "October 9, 2026",
                "recipient": "Alex Example",
                "address": "100 Example Street\nExample City, ZZ 00000",
                "subject": "Fictional matter",
                "signoff": "Example Counsel",
                "author": "Authorized reviewer",
            })()
            creator.create(args)
            document = Document(output)
            header = "\n".join(
                paragraph.text
                for section in document.sections
                for paragraph in section.header.paragraphs
            )
            self.assertIn("Northstar Legal Lab", header)
            self.assertNotIn("ORGANIZATION NAME", header)

    def test_installer_excludes_build_residue(self):
        installer = load_script("install_private")
        ignored = installer.ignored("unused", ["a.pyc", "pkg.egg-info", "README.md"])
        self.assertEqual(ignored, {"a.pyc", "pkg.egg-info"})

    def test_documented_local_references_exist(self):
        markdown_files = [ROOT / "README.md"]
        markdown_files.extend((ROOT / "docs").glob("*.md"))
        markdown_files.extend((ROOT / "skills").glob("*/SKILL.md"))

        missing = []
        for path in markdown_files:
            text = path.read_text(encoding="utf-8")
            for raw_target in re.findall(r"\[[^]]+\]\(([^)]+)\)", text):
                target = unquote(raw_target.split("#", 1)[0])
                if not target or "://" in target:
                    continue
                if not (path.parent / target).resolve().is_file():
                    missing.append(f"{path.relative_to(ROOT)} -> {raw_target}")
            for raw_target in re.findall(r"\$\{LAWYEREDOS_ROOT\}/([A-Za-z0-9_./-]+)", text):
                target = raw_target.rstrip(".,;:")
                if not (ROOT / target).exists():
                    missing.append(f"{path.relative_to(ROOT)} -> {target}")
        self.assertEqual(missing, [])

    def test_documented_clean_install_inputs_are_bundled(self):
        required = [
            "LICENSE",
            "LICENSE-DECISION.md",
            "SETUP.md",
            "requirements-lock.txt",
            "requirements-test.txt",
            "config.example.toml",
            "scripts/setup_demo.py",
            "scripts/setup_wizard.py",
            "scripts/smoke_setup_wizard.py",
            "scripts/validate_setup.py",
            "docs/WIZARD.md",
            ".claude/skills/lawyeredos-setup/SKILL.md",
            "assets/neutral-letterhead.docx",
            "vendor/legal-redline-tools/LICENSE",
        ]
        self.assertEqual([path for path in required if not (ROOT / path).is_file()], [])
        lock = (ROOT / "requirements-lock.txt").read_text(encoding="utf-8")
        self.assertRegex(lock, r"(?m)^setuptools==")
        self.assertRegex(lock, r"(?m)^wheel==")
        tests = (ROOT / "requirements-test.txt").read_text(encoding="utf-8")
        self.assertRegex(tests, r"(?m)^pytest==")

    def test_plain_language_setup_entrypoint_is_documented(self):
        prompt = (
            "Set up LawyeredOS for me. Follow the setup instructions in this folder, "
            "ask me what you need, and check with me before changing existing files."
        )
        self.assertIn(prompt, (ROOT / "SETUP.md").read_text(encoding="utf-8"))
        self.assertIn(prompt, (ROOT / "README.md").read_text(encoding="utf-8"))

    def test_public_contribution_surfaces_are_dual_host_and_least_privilege(self):
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("Claude Desktop Code Local", readme)
        self.assertIn("Codex local project", readme)
        self.assertIn("0f3eccf6e62e0bcb7ffbb88a6c14dd2f96c8a05acba73b93bb4bff3173af7cd5", readme)
        self.assertIn("b468300f81bfdcc88398964a2808f6484c5cd2fb56b44d48e6721ab2d7d6e366", readme)

        pull_request = (ROOT / ".github" / "pull_request_template.md").read_text(encoding="utf-8")
        self.assertIn("no authentic client data", pull_request)
        self.assertIn("published ZIP", pull_request)

        workflow = (ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8")
        self.assertIn("contents: read", workflow)
        self.assertIn("github.event.repository.private == false", workflow)
        self.assertIn("persist-credentials: false", workflow)
        self.assertNotIn("pull_request_target", workflow)
        self.assertNotIn("secrets.", workflow)

        security = (ROOT / "SECURITY.md").read_text(encoding="utf-8")
        self.assertIn("/security/advisories/new", security)
        self.assertIn("Report a vulnerability", " ".join(security.split()))

        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("Screen the local register for potential conflicts", readme)
        self.assertIn("no dedicated clearance-status field", readme)
        tracker_skill = (
            ROOT / "skills" / "practice-matter-tracker" / "SKILL.md"
        ).read_text(encoding="utf-8")
        self.assertIn("--status all", tracker_skill)
        self.assertIn("final conflict decision", tracker_skill)


if __name__ == "__main__":
    unittest.main()

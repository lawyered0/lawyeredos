from __future__ import annotations

import base64
import io
import json
import sqlite3
import subprocess
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from lawyeredos.integrations import GmailReadonlyAdapter
from lawyeredos.imap_mail import GmailImapReadonlyAdapter, credentials_from_config
from lawyeredos.content_review import REVIEW_ACK_SCHEMA
from lawyeredos.mail_intake import (
    acknowledge_review_chunk,
    main,
    next_review_chunk,
    plan,
    process_bundle,
)
from lawyeredos.tracker_cli import initialize


class MailIntakeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workspace = self.root / "workspace"
        self.matter = self.workspace / "2032-001-northstar"
        self.matter.mkdir(parents=True)
        (self.matter / "_matter-brief.md").write_text("# Fictional matter\n", encoding="utf-8")
        self.database = self.workspace / "matters.sqlite3"
        initialize(self.database)
        with sqlite3.connect(self.database) as connection:
            connection.execute(
                """INSERT INTO matters(file_number, client, description, matter_folder,
                status, opposing_party, primary_email, created_at, updated_at)
                VALUES (?, ?, ?, ?, 'open', ?, ?, ?, ?)""",
                (
                    "2032-001", "Northstar Bakery Ltd.", "Fictional supply dispute",
                    "2032-001-northstar", "Harbor Paper Co.",
                    "client@example.invalid", "2032-01-01", "2032-01-01",
                ),
            )
        self.template = self.root / "template.docx"
        self.template.write_bytes(b"fixture")
        self.config = self.root / "config.toml"
        self.write_config("fixture-json")

    def tearDown(self):
        self.temp.cleanup()

    def write_config(self, adapter: str) -> None:
        self.config.write_text(f'''schema_version = 1
[workspace]
root = "{self.workspace}"
scratch = "{self.workspace / '.lawyeredos' / 'scratch'}"
timezone = "UTC"
[tracker]
database = "{self.database}"
[documents]
template = "{self.template}"
author = "Authorized reviewer"
[jurisdiction]
profile = "none"
[integrations]
mail_adapter = "{adapter}"
calendar_adapter = "unavailable"
practice_manager_adapter = "unavailable"
[calendar]
calendar_id = ""
[review]
external_ai = false
''', encoding="utf-8")

    def message(self, message_id: str = "m1", *, attachment: dict | None = None) -> dict:
        return {
            "id": message_id, "thread_id": "thread-1", "received_at": "2032-02-01T12:00:00+00:00",
            "from": "client@example.invalid", "to": "lawyer@example.invalid", "cc": "",
            "subject": "Fictional supply dispute", "snippet": "Please review the attached note.",
            "body": "Northstar Bakery asks counsel to review the attached fictional note.",
            "body_complete": True, "attachments": [] if attachment is None else [attachment],
        }

    def attachment(self, content: bytes = b"Fictional attachment text.\n", **changes) -> dict:
        value = {
            "id": "a1", "filename": "../client-note.txt", "mime_type": "text/plain",
            "size": len(content), "content_base64": base64.b64encode(content).decode("ascii"),
        }
        value.update(changes)
        return value

    def bundle(self, planned: dict, messages: list[dict], *, complete: bool = True) -> dict:
        return {
            "schema": "lawyeredos-mail-bundle-v1", "source": "fixture",
            "query": planned["query"], "messages": messages,
            "coverage": {"complete": complete, "completed_at": "2032-02-01T13:00:00+00:00"},
        }

    def acknowledge_all(self, identifier: str = "Northstar", limit: int | None = None) -> list[dict]:
        acknowledgements = []
        while limit is None or len(acknowledgements) < limit:
            pending = next_review_chunk(self.config, identifier)
            if pending["status"] == "no_pending_review":
                break
            acknowledgement = {
                "schema": REVIEW_ACK_SCHEMA,
                "reviewed_text_sha256": pending["text_sha256"],
                "reviewer": "fictional-test-reviewer",
                "summary": f"Reviewed bounded chunk {pending['chunk_index']} as fictional evidence.",
                "evidence": [{
                    "locator": json.dumps(pending["locators"][0], sort_keys=True),
                    "observation": "The chunk was read as untrusted fictional matter evidence.",
                }],
                "limitations": [],
            }
            acknowledgements.append(
                acknowledge_review_chunk(
                    self.config, identifier, pending["review_id"], pending["chunk_index"],
                    pending["text_sha256"], acknowledgement,
                )
            )
        return acknowledgements

    def test_complete_repeat_and_new_message_flow(self):
        first_plan = plan(self.config, "Northstar")
        result = process_bundle(
            self.config, "Northstar",
            self.bundle(first_plan, [self.message(attachment=self.attachment())]),
        )
        self.assertEqual(result["status"], "partial")
        self.assertFalse(result["checkpoint_advanced"])
        self.assertEqual(result["messages"][0]["status"], "partial")
        self.assertEqual(len(self.acknowledge_all()), 2)
        reviewed_result = process_bundle(
            self.config, "Northstar",
            self.bundle(first_plan, [self.message(attachment=self.attachment())]),
        )
        self.assertEqual(reviewed_result["status"], "complete")
        self.assertTrue(reviewed_result["checkpoint_advanced"])
        reviewed = reviewed_result["messages"][0]
        self.assertEqual(reviewed["status"], "reviewed")
        self.assertTrue(reviewed["untrusted_evidence"])
        saved = Path(reviewed["attachments"][0]["saved_path"])
        self.assertEqual(saved.name, "client-note.txt")
        self.assertEqual(saved.read_bytes(), b"Fictional attachment text.\n")

        second_plan = plan(self.config, "2032-001")
        self.assertIn("newer_than:30d", second_plan["query"])
        repeated = process_bundle(
            self.config, "2032-001",
            self.bundle(second_plan, [self.message(attachment=self.attachment())]),
        )
        self.assertEqual(repeated["messages"], [{"message_id": "m1", "status": "already_reviewed"}])
        self.assertEqual(len(list((self.matter / "email-attachments").iterdir())), 1)

        third_plan = plan(self.config, "Northstar")
        new_message = self.message("m2")
        new_message["thread_id"] = "thread-2"
        new_message["body"] = "A new Northstar update with no attachment."
        newest = process_bundle(self.config, "Northstar", self.bundle(third_plan, [new_message]))
        self.assertEqual(newest["messages"][0]["status"], "partial")
        self.assertEqual(len(self.acknowledge_all()), 1)
        newest = process_bundle(self.config, "Northstar", self.bundle(third_plan, [new_message]))
        self.assertEqual(newest["messages"][0]["status"], "reviewed")

    def test_partial_attachment_retry_does_not_advance_checkpoint(self):
        first_plan = plan(self.config, "Northstar")
        missing = self.attachment()
        missing.pop("content_base64")
        partial = process_bundle(
            self.config, "Northstar", self.bundle(first_plan, [self.message(attachment=missing)]),
        )
        self.assertEqual(partial["status"], "partial")
        self.assertFalse(partial["checkpoint_advanced"])
        self.assertEqual(partial["messages"][0]["attachments"][0]["status"], "deferred")
        self.assertIsNone(plan(self.config, "Northstar")["checkpoint"])
        self.assertEqual(len(self.acknowledge_all()), 1)

        retry_plan = plan(self.config, "Northstar")
        retried = process_bundle(
            self.config, "Northstar",
            self.bundle(retry_plan, [self.message(attachment=self.attachment())]),
        )
        self.assertEqual(retried["status"], "partial")
        self.assertEqual(len(self.acknowledge_all()), 1)
        retried = process_bundle(
            self.config, "Northstar", self.bundle(retry_plan, [self.message(attachment=self.attachment())]),
        )
        self.assertEqual(retried["status"], "complete")
        self.assertIsNotNone(plan(self.config, "Northstar")["checkpoint"])

    def test_docx_is_reviewed_and_existing_different_file_is_not_overwritten(self):
        from docx import Document

        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Northstar fictional document evidence.")
        document.save(buffer)
        existing = self.matter / "email-attachments"
        existing.mkdir()
        (existing / "client-note.docx").write_bytes(b"pre-existing user file")

        planned = plan(self.config, "Northstar")
        attachment = self.attachment(
            buffer.getvalue(), filename="client-note.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        result = process_bundle(
            self.config, "Northstar", self.bundle(planned, [self.message(attachment=attachment)]),
        )
        reviewed = result["messages"][0]["attachments"][0]
        self.assertEqual(reviewed["status"], "review_required")
        self.assertNotEqual(Path(reviewed["saved_path"]).name, "client-note.docx")
        self.assertEqual((existing / "client-note.docx").read_bytes(), b"pre-existing user file")
        self.assertEqual(len(self.acknowledge_all()), 2)
        completed = process_bundle(
            self.config, "Northstar", self.bundle(planned, [self.message(attachment=attachment)]),
        )
        self.assertEqual(completed["messages"][0]["attachments"][0]["status"], "reviewed")

    def test_incomplete_coverage_holds_checkpoint_but_dedupes_review(self):
        first_plan = plan(self.config, "Northstar")
        partial = process_bundle(
            self.config, "Northstar", self.bundle(first_plan, [self.message()], complete=False),
        )
        self.assertEqual(partial["messages"][0]["status"], "partial")
        self.assertFalse(partial["checkpoint_advanced"])
        self.assertIsNone(plan(self.config, "Northstar")["checkpoint"])
        self.assertEqual(len(self.acknowledge_all()), 1)

        partial = process_bundle(
            self.config, "Northstar", self.bundle(first_plan, [self.message()], complete=False),
        )
        self.assertEqual(partial["messages"][0]["status"], "reviewed")
        self.assertFalse(partial["checkpoint_advanced"])

        retry_plan = plan(self.config, "Northstar")
        completed = process_bundle(
            self.config, "Northstar", self.bundle(retry_plan, [self.message()], complete=True),
        )
        self.assertEqual(completed["messages"][0]["status"], "already_reviewed")
        self.assertTrue(completed["checkpoint_advanced"])

    def test_incomplete_body_and_missing_attachment_bytes_remain_retryable(self):
        planned = plan(self.config, "Northstar")
        message = self.message(attachment=self.attachment())
        message["body_complete"] = False
        message["attachments"][0].pop("content_base64")
        result = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(result["status"], "partial")
        self.assertIn("body retrieval was incomplete", result["messages"][0]["problems"][0])
        self.assertEqual(result["messages"][0]["attachments"][0]["status"], "deferred")
        self.assertIsNone(plan(self.config, "Northstar")["checkpoint"])

    def test_long_body_and_attachment_review_resume_without_early_checkpoint(self):
        planned = plan(self.config, "Northstar")
        message = self.message(attachment=self.attachment(b"x" * 12_001))
        message["body"] = "Fictional supply dispute " + ("y" * 12_001)
        result = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(result["status"], "partial")
        self.assertFalse(result["checkpoint_advanced"])
        body_review = result["messages"][0]["body_review"]
        attachment_review = result["messages"][0]["attachments"][0]["review"]
        self.assertGreater(body_review["chunks_total"], 1)
        self.assertGreater(attachment_review["chunks_total"], 1)
        self.assertEqual(self.acknowledge_all(limit=1)[0]["remaining_chunks"], 1)
        interrupted = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertFalse(interrupted["checkpoint_advanced"])
        self.assertEqual(interrupted["messages"][0]["body_review"]["chunks_reviewed"], 1)
        self.assertGreaterEqual(len(self.acknowledge_all()), 3)
        completed = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(completed["status"], "complete")
        self.assertTrue(completed["checkpoint_advanced"])

    def test_long_docx_review_is_paragraph_aware_and_resumable(self):
        from docx import Document

        document = Document()
        for number in range(1, 5):
            document.add_paragraph(f"Paragraph {number}: " + (chr(64 + number) * 7_000))
        buffer = io.BytesIO()
        document.save(buffer)
        planned = plan(self.config, "Northstar")
        attachment = self.attachment(
            buffer.getvalue(), filename="long-agreement.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        message = self.message(attachment=attachment)
        first = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        review = first["messages"][0]["attachments"][0]["review"]
        self.assertGreater(review["chunks_total"], 2)
        self.assertEqual(review["coverage"]["details"]["reviewed_parts"], ["word/document.xml"])
        self.assertFalse(first["checkpoint_advanced"])

        self.assertEqual(len(self.acknowledge_all(limit=2)), 2)
        resumed = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        resumed_review = resumed["messages"][0]["attachments"][0]["review"]
        self.assertGreater(resumed_review["chunks_reviewed"], 0)
        self.assertLess(resumed_review["chunks_reviewed"], resumed_review["chunks_total"])
        self.assertFalse(resumed["checkpoint_advanced"])

        self.acknowledge_all()
        completed = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(completed["status"], "complete")
        self.assertTrue(completed["checkpoint_advanced"])

    def test_long_pdf_review_is_page_aware_and_resumable(self):
        planned = plan(self.config, "Northstar")
        pdf = self.attachment(
            b"%PDF-1.4\nfictional bounded test\n", filename="long-record.pdf",
            mime_type="application/pdf",
        )
        message = self.message(attachment=pdf)

        def fake_extract(_source, output):
            output.write_text(
                ("Page one evidence. " * 800) + "\f" +
                ("Page two evidence. " * 800) + "\f" +
                ("Page three evidence. " * 800),
                encoding="utf-8",
            )
            return {"ok": True, "needs_ocr": False}

        with patch("lawyeredos.content_review.extract_pdf", side_effect=fake_extract):
            first = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        review = first["messages"][0]["attachments"][0]["review"]
        self.assertEqual(review["coverage"]["details"]["pages"], 3)
        self.assertGreaterEqual(review["chunks_total"], 3)
        locators = []
        while True:
            pending = next_review_chunk(self.config, "Northstar")
            if pending["status"] == "no_pending_review":
                break
            locators.extend(pending["locators"])
            self.acknowledge_all(limit=1)
        self.assertEqual({item["page"] for item in locators if item["kind"] == "pdf"}, {1, 2, 3})
        with patch("lawyeredos.content_review.extract_pdf", side_effect=fake_extract):
            completed = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(completed["status"], "complete")

    def test_xlsx_review_covers_every_sheet_including_formulas_and_hidden_sheet(self):
        from openpyxl import Workbook

        workbook = Workbook()
        first_sheet = workbook.active
        first_sheet.title = "Summary"
        first_sheet.append(["Client", "Amount", "Formula"])
        first_sheet.append(["Northstar Bakery", 1250, "=B2*2"])
        second_sheet = workbook.create_sheet("Details")
        second_sheet.append(["Harbor Paper", "Fictional supply dispute"])
        hidden_sheet = workbook.create_sheet("Hidden assumptions")
        hidden_sheet.sheet_state = "hidden"
        hidden_sheet.append(["Assumption", "Not independently verified"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        planned = plan(self.config, "Northstar")
        attachment = self.attachment(
            buffer.getvalue(), filename="damages.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        message = self.message(attachment=attachment)
        first = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        review = first["messages"][0]["attachments"][0]["review"]
        sheets = review["coverage"]["details"]["sheets"]
        self.assertEqual([item["title"] for item in sheets], ["Summary", "Details", "Hidden assumptions"])
        self.assertEqual(sheets[-1]["state"], "hidden")
        pending_text = []
        while next_review_chunk(self.config, "Northstar")["status"] == "review_required":
            pending = next_review_chunk(self.config, "Northstar")
            pending_text.append(pending["text"])
            self.acknowledge_all(limit=1)
        self.assertIn("C2==B2*2", "\n".join(pending_text))
        completed = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(completed["status"], "complete")

    def test_xlsx_visual_objects_hold_completion(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Month", "Amount"])
        sheet.append(["One", 10])
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        sheet.add_chart(chart, "D2")
        buffer = io.BytesIO()
        workbook.save(buffer)
        planned = plan(self.config, "Northstar")
        attachment = self.attachment(
            buffer.getvalue(), filename="visual-ledger.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        result = process_bundle(
            self.config, "Northstar", self.bundle(planned, [self.message(attachment=attachment)]),
        )
        reviewed = result["messages"][0]["attachments"][0]
        self.assertEqual(reviewed["status"], "extraction_incomplete")
        self.assertEqual(reviewed["review"]["status"], "blocked")
        self.assertEqual(reviewed["review"]["coverage"]["details"]["visual_objects"], 1)
        self.assertFalse(result["checkpoint_advanced"])

    def test_extracted_content_cap_blocks_giant_review_item(self):
        planned = plan(self.config, "Northstar")
        attachment = self.attachment(b"bounded evidence " * 20)
        with patch("lawyeredos.content_review.MAX_REVIEW_CONTENT_CHARS", 100):
            result = process_bundle(
                self.config, "Northstar", self.bundle(planned, [self.message(attachment=attachment)]),
            )
        reviewed = result["messages"][0]["attachments"][0]
        self.assertEqual(reviewed["status"], "extraction_incomplete")
        self.assertIn("exceeds", reviewed["review"]["coverage"]["reason"])
        self.assertEqual(reviewed["review"]["chunks_total"], 0)
        self.assertFalse(result["checkpoint_advanced"])

    def test_review_acknowledgement_is_hash_bound_and_cannot_advance_early(self):
        planned = plan(self.config, "Northstar")
        message = self.message(attachment=self.attachment(b"Evidence " * 2_000))
        first = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        pending = next_review_chunk(self.config, "Northstar")
        acknowledgement = {
            "schema": REVIEW_ACK_SCHEMA,
            "reviewed_text_sha256": "0" * 64,
            "reviewer": "fictional-test-reviewer",
            "summary": "This acknowledgement deliberately has the wrong hash.",
            "evidence": [{"locator": "line 1", "observation": "Fictional evidence."}],
            "limitations": [],
        }
        with self.assertRaisesRegex(ValueError, "hash"):
            acknowledge_review_chunk(
                self.config, "Northstar", pending["review_id"], pending["chunk_index"],
                pending["text_sha256"], acknowledgement,
            )
        resumed = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(resumed["messages"][0]["body_review"]["chunks_reviewed"], 0)
        self.assertFalse(resumed["checkpoint_advanced"])
        self.assertEqual(first["messages"][0]["body_review"]["review_id"], pending["review_id"])

    def test_query_result_without_exact_matter_pair_is_held(self):
        planned = plan(self.config, "Northstar")
        message = self.message()
        message["from"] = "unrelated@example.invalid"
        message["subject"] = "Northstar Bakery Ltd."
        message["body"] = "A party-only match without a matter discriminator."
        result = process_bundle(self.config, "Northstar", self.bundle(planned, [message]))
        self.assertEqual(result["messages"][0]["status"], "held_ambiguous")
        self.assertFalse(result["checkpoint_advanced"])

    def test_bundle_source_query_and_timestamp_are_verified(self):
        planned = plan(self.config, "Northstar")
        wrong_source = self.bundle(planned, [self.message()])
        wrong_source["source"] = "codex-session-tool"
        with self.assertRaisesRegex(ValueError, "source"):
            process_bundle(self.config, "Northstar", wrong_source)

        wrong_query = self.bundle(planned, [self.message()])
        wrong_query["query"] = "newer_than:999d"
        with self.assertRaisesRegex(ValueError, "query"):
            process_bundle(self.config, "Northstar", wrong_query)

        bad_timestamp = self.message()
        bad_timestamp["received_at"] = "2032-02-01 12:00:00"
        with self.assertRaisesRegex(ValueError, "UTC offset"):
            process_bundle(self.config, "Northstar", self.bundle(planned, [bad_timestamp]))

    def test_checkpoints_are_scoped_to_the_configured_mail_source(self):
        fixture_plan = plan(self.config, "Northstar")
        result = process_bundle(
            self.config, "Northstar", self.bundle(fixture_plan, [self.message()]),
        )
        self.assertFalse(result["checkpoint_advanced"])
        self.acknowledge_all()
        result = process_bundle(
            self.config, "Northstar", self.bundle(fixture_plan, [self.message()]),
        )
        self.assertTrue(result["checkpoint_advanced"])

        self.write_config("codex-session-tools")
        codex_plan = plan(self.config, "Northstar")
        self.assertEqual(codex_plan["source"], "codex-session-tool")
        self.assertIsNone(codex_plan["checkpoint"])
        self.assertIn("newer_than:30d", codex_plan["query"])

        self.write_config("unavailable")
        unavailable_plan = plan(self.config, "Northstar")
        self.assertIsNone(unavailable_plan["source"])
        self.assertIsNone(unavailable_plan["checkpoint"])

    def test_unsafe_oversized_and_executable_attachments_are_deferred(self):
        planned = plan(self.config, "Northstar")
        executable = self.attachment(
            b"MZfake", filename="invoice.exe", mime_type="application/octet-stream",
        )
        result = process_bundle(
            self.config, "Northstar", self.bundle(planned, [self.message(attachment=executable)]),
        )
        self.assertEqual(result["status"], "partial")
        self.assertIn("blocked", result["messages"][0]["attachments"][0]["error"])

        with patch("lawyeredos.mail_intake.MAX_ATTACHMENT_BYTES", 3):
            planned = plan(self.config, "Northstar")
            oversized = process_bundle(
                self.config, "Northstar",
                self.bundle(planned, [self.message(attachment=self.attachment(b"four"))]),
            )
        self.assertEqual(oversized["status"], "partial")
        self.assertIn("exceeds", oversized["messages"][0]["attachments"][0]["error"])

        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as output:
            output.writestr("word/document.xml", "<document/>")
            output.writestr("word/embeddings/oleObject1.bin", b"payload")
        planned = plan(self.config, "Northstar")
        active_docx = self.attachment(
            archive.getvalue(), filename="active.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        active = process_bundle(
            self.config, "Northstar", self.bundle(planned, [self.message(attachment=active_docx)]),
        )
        self.assertIn("active content", active["messages"][0]["attachments"][0]["error"])

    def test_managed_private_session_bundle_is_deleted_after_processing(self):
        planned = plan(self.config, "Northstar")
        folder = self.workspace / ".lawyeredos" / "scratch" / "mail-intake-session"
        folder.mkdir(parents=True)
        folder.chmod(0o700)
        bundle_path = folder / "one-run.json"
        bundle_path.write_text(json.dumps(self.bundle(planned, [self.message()])), encoding="utf-8")
        bundle_path.chmod(0o600)
        output = io.StringIO()
        with redirect_stdout(output):
            status = main([
                "process", "--config", str(self.config), "--identifier", "Northstar",
                "--bundle", str(bundle_path), "--delete-bundle",
            ])
        self.assertEqual(status, 0)
        self.assertFalse(bundle_path.exists())

    def test_managed_review_ack_is_hash_bound_and_deleted_after_commit(self):
        planned = plan(self.config, "Northstar")
        process_bundle(self.config, "Northstar", self.bundle(planned, [self.message()]))
        pending = next_review_chunk(self.config, "Northstar")
        folder = self.workspace / ".lawyeredos" / "scratch" / "mail-review-acks"
        folder.mkdir(parents=True)
        folder.chmod(0o700)
        ack_path = folder / "ack.json"
        ack_path.write_text(json.dumps({
            "schema": REVIEW_ACK_SCHEMA,
            "reviewed_text_sha256": pending["text_sha256"],
            "reviewer": "fictional-test-reviewer",
            "summary": "The complete fictional message chunk was substantively reviewed.",
            "evidence": [{
                "locator": "message body",
                "observation": "It requests review of the fictional attachment.",
            }],
            "limitations": [],
        }), encoding="utf-8")
        ack_path.chmod(0o600)
        output = io.StringIO()
        with redirect_stdout(output):
            status = main([
                "review-ack", "--config", str(self.config), "--identifier", "Northstar",
                "--review-id", pending["review_id"], "--chunk-index", str(pending["chunk_index"]),
                "--text-sha256", pending["text_sha256"], "--ack-file", str(ack_path),
                "--delete-ack",
            ])
        self.assertEqual(status, 0)
        self.assertFalse(ack_path.exists())

    def test_legacy_unacknowledged_completion_state_is_invalidated(self):
        planned = plan(self.config, "Northstar")
        process_bundle(self.config, "Northstar", self.bundle(planned, [self.message()]))
        state = self.workspace / ".lawyeredos" / "data" / "mail-intake.sqlite3"
        with sqlite3.connect(state) as connection:
            connection.execute(
                "UPDATE state_meta SET value='legacy-extraction-counted-as-review' "
                "WHERE key='review_protocol'"
            )
            connection.execute(
                "INSERT INTO scan_state VALUES (?, ?, ?, ?, ?)",
                ("2032-001", "fixture", "2032-02-01T13:00:00+00:00", "fixture", planned["query"]),
            )
        refreshed = plan(self.config, "Northstar")
        self.assertIsNone(refreshed["checkpoint"])
        with sqlite3.connect(state) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM scan_state").fetchone()[0], 0)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM content_review_items").fetchone()[0], 0)
        self.assertEqual(state.stat().st_mode & 0o777, 0o600)
        self.assertEqual(state.parent.stat().st_mode & 0o777, 0o700)

    def test_symlink_destination_and_ambiguous_matter_fail_closed(self):
        outside = self.root / "outside"
        outside.mkdir()
        (self.matter / "email-attachments").symlink_to(outside, target_is_directory=True)
        planned = plan(self.config, "2032-001")
        with self.assertRaisesRegex(ValueError, "symlink|unsafe"):
            process_bundle(self.config, "2032-001", self.bundle(planned, [self.message()]))
        (self.matter / "email-attachments").unlink()

        second = self.workspace / "2032-002-northstar"
        second.mkdir()
        with sqlite3.connect(self.database) as connection:
            connection.execute(
                """INSERT INTO matters(file_number, client, description, matter_folder,
                status, created_at, updated_at) VALUES
                ('2032-002', 'Northstar Holdings Ltd.', 'Other fictional matter',
                '2032-002-northstar', 'open', '2032-01-01', '2032-01-01')"""
            )
        ambiguous = plan(self.config, "Northstar")
        self.assertFalse(ambiguous["ok"])
        self.assertEqual(ambiguous["status"], "ambiguous")

    def test_unavailable_adapter_reports_without_creating_state(self):
        self.write_config("unavailable")
        result = plan(self.config, "Northstar")
        self.assertEqual(result["status"], "mail_unavailable")
        self.assertEqual(result["route"], "unavailable")
        self.assertFalse((self.workspace / ".lawyeredos" / "data" / "mail-intake.sqlite3").exists())

    def test_imap_provider_bundle_processes_and_deduplicates_end_to_end(self):
        self.write_config("gmail-imap-readonly")
        client = FakeImapClient({101: (int("1a2b3c", 16), int("4d5e6f", 16), imap_message())})
        adapter = GmailImapReadonlyAdapter(
            "fictional@example.invalid", "app-password-fixture", lambda _host: client
        )
        planned = plan(self.config, "2032-001")
        first = process_bundle(
            self.config, "2032-001", adapter.intake(planned["query"], 5)
        )
        self.assertEqual(first["status"], "partial")
        self.assertEqual(first["messages"][0]["status"], "partial")
        self.assertEqual(first["messages"][0]["attachments"][0]["status"], "review_required")
        self.assertEqual(len(self.acknowledge_all("2032-001")), 2)
        completed_client = FakeImapClient(
            {101: (int("1a2b3c", 16), int("4d5e6f", 16), imap_message())}
        )
        completed = process_bundle(
            self.config, "2032-001",
            GmailImapReadonlyAdapter(
                "fictional@example.invalid", "app-password-fixture", lambda _host: completed_client
            ).intake(plan(self.config, "2032-001")["query"], 5),
        )
        self.assertEqual(completed["status"], "complete")
        self.assertEqual(completed["messages"][0]["attachments"][0]["status"], "reviewed")
        self.assertEqual(len(list((self.matter / "email-attachments").iterdir())), 1)

        repeated_client = FakeImapClient(
            {101: (int("1a2b3c", 16), int("4d5e6f", 16), imap_message())}
        )
        repeated_adapter = GmailImapReadonlyAdapter(
            "fictional@example.invalid", "app-password-fixture", lambda _host: repeated_client
        )
        repeated = process_bundle(
            self.config, "Northstar",
            repeated_adapter.intake(plan(self.config, "Northstar")["query"], 5),
        )
        self.assertEqual(
            repeated["messages"], [{"message_id": "1a2b3c", "status": "already_reviewed"}]
        )
        self.assertEqual(len(list((self.matter / "email-attachments").iterdir())), 1)


class GmailIntakeTransportTests(unittest.TestCase):
    def test_full_message_and_attachment_use_read_only_fake_transport(self):
        calls = []
        body = base64.urlsafe_b64encode(b"Northstar body").decode("ascii").rstrip("=")
        attachment = base64.urlsafe_b64encode(b"fictional text").decode("ascii").rstrip("=")

        def transport(method, url, headers, request_body):
            calls.append((method, url, request_body))
            if url.endswith("/profile"):
                return 200, b'{"emailAddress":"fictional@example.invalid"}'
            if "/attachments/" in url:
                return 200, json.dumps({"data": attachment}).encode()
            if "format=full" in url:
                return 200, json.dumps({
                    "id": "m1", "threadId": "t1", "internalDate": "1950000000000",
                    "snippet": "Northstar", "payload": {
                        "headers": [
                            {"name": "From", "value": "client@example.invalid"},
                            {"name": "Subject", "value": "Northstar update"},
                        ],
                        "parts": [
                            {"mimeType": "text/plain", "body": {"data": body}},
                            {"mimeType": "text/plain", "filename": "note.txt",
                             "body": {"attachmentId": "a1", "size": 14}},
                        ],
                    },
                }).encode()
            return 200, b'{"messages":[{"id":"m1"}]}'

        result = GmailReadonlyAdapter("fake-token", transport).intake("Northstar newer_than:30d", 5)
        self.assertTrue(result["coverage"]["complete"])
        self.assertEqual(result["messages"][0]["body"], "Northstar body")
        self.assertEqual(
            base64.b64decode(result["messages"][0]["attachments"][0]["content_base64"]),
            b"fictional text",
        )
        self.assertTrue(all(method == "GET" and request_body is None for method, _url, request_body in calls))

    def test_oversized_body_is_marked_incomplete(self):
        body = base64.urlsafe_b64encode(b"Northstar body").decode("ascii").rstrip("=")

        def transport(method, url, headers, request_body):
            if url.endswith("/profile"):
                return 200, b'{"emailAddress":"fictional@example.invalid"}'
            if "format=full" in url:
                return 200, json.dumps({
                    "id": "m1", "threadId": "t1", "internalDate": "1950000000000",
                    "payload": {"headers": [], "mimeType": "text/plain", "body": {"data": body}},
                }).encode()
            return 200, b'{"messages":[{"id":"m1"}]}'

        with patch("lawyeredos.integrations.MAX_MAIL_BODY_CHARS", 5):
            result = GmailReadonlyAdapter("fake-token", transport).intake("Northstar", 5)
        self.assertEqual(result["messages"][0]["body"], "North")
        self.assertFalse(result["messages"][0]["body_complete"])

    def test_attachment_backed_body_inline_image_and_aggregate_limit_are_explicit(self):
        encoded_body = base64.urlsafe_b64encode(b"Northstar fetched body").decode("ascii").rstrip("=")
        encoded_image = base64.urlsafe_b64encode(b"image-bytes").decode("ascii").rstrip("=")
        calls = []

        def transport(method, url, headers, request_body):
            calls.append(url)
            if url.endswith("/profile"):
                return 200, b'{"emailAddress":"fictional@example.invalid"}'
            if "/attachments/body-1" in url:
                return 200, json.dumps({"data": encoded_body}).encode()
            if "/attachments/image-1" in url:
                return 200, json.dumps({"data": encoded_image}).encode()
            if "format=full" in url:
                return 200, json.dumps({
                    "id": "m1", "threadId": "t1", "internalDate": "1950000000000",
                    "payload": {"headers": [], "parts": [
                        {"mimeType": "text/plain", "body": {"attachmentId": "body-1", "size": 22}},
                        {"mimeType": "image/png", "filename": "signature.png",
                         "headers": [{"name": "Content-Disposition", "value": "inline"}],
                         "body": {"attachmentId": "image-1", "size": 11}},
                    ]},
                }).encode()
            return 200, b'{"messages":[{"id":"m1"}]}'

        with patch("lawyeredos.integrations.MAX_MAIL_RUN_ATTACHMENT_BYTES", 5):
            result = GmailReadonlyAdapter("fake-token", transport).intake("Northstar", 5)
        message = result["messages"][0]
        self.assertEqual(message["body"], "Northstar fetched body")
        self.assertTrue(message["body_complete"])
        self.assertEqual(message["attachments"][0]["filename"], "signature.png")
        self.assertTrue(message["attachments"][0]["inline"])
        self.assertIn("bounded", message["attachments"][0]["retrieval_error"])
        self.assertFalse(any("/attachments/image-1" in url for url in calls))


class FakeImapClient:
    def __init__(self, messages: dict[int, tuple[int, int, bytes]]):
        self.messages = messages
        self.calls = []
        self.selected = None

    def login(self, account, password):
        self.calls.append(("LOGIN", account, bool(password)))
        return "OK", [b"authenticated"]

    def list(self):
        self.calls.append(("LIST",))
        return "OK", [
            b'(\\HasNoChildren) "/" "INBOX"',
            b'(\\HasNoChildren \\All) "/" "[Gmail]/Todo el correo"',
        ]

    def select(self, mailbox, readonly=False):
        self.calls.append(("SELECT", mailbox, readonly))
        self.selected = (mailbox, readonly)
        return "OK", [str(len(self.messages)).encode()]

    def uid(self, command, *args):
        self.calls.append(("UID", command, *args))
        if command == "SEARCH":
            return "OK", [" ".join(str(uid) for uid in self.messages).encode()]
        if command != "FETCH":
            return "BAD", []
        uid = int(args[0])
        gm_message_id, gm_thread_id, raw = self.messages[uid]
        query = args[1]
        if "HEADER.FIELDS" in query:
            header = raw.split(b"\n\n", 1)[0] + b"\n\n"
            attributes = (
                f'{uid} (UID {uid} X-GM-MSGID {gm_message_id} '
                f'X-GM-THRID {gm_thread_id} INTERNALDATE "01-Feb-2032 12:00:00 +0000" '
                f'RFC822.SIZE {len(raw)} BODY[HEADER.FIELDS] {{{len(header)}}}'
            ).encode()
            return "OK", [(attributes, header), b")"]
        if "BODY.PEEK[]" in query:
            return "OK", [(f"{uid} (UID {uid} BODY[] {{{len(raw)}}}".encode(), raw), b")"]
        return "BAD", []

    def logout(self):
        self.calls.append(("LOGOUT",))
        return "BYE", [b"closed"]


def imap_message(*, include_complex=False) -> bytes:
    message = EmailMessage()
    message["From"] = "Fictional Client <client@example.invalid>"
    message["To"] = "Fictional Lawyer <lawyer@example.invalid>"
    message["Subject"] = "Fictional supply dispute"
    message.set_content("Northstar Bakery complete fictional IMAP message body.")
    message.add_attachment(
        b"Northstar fictional attachment.\n",
        maintype="text", subtype="plain", filename="evidence.txt",
    )
    if include_complex:
        message.add_attachment(
            b"%PDF-1.4\nfictional unnamed\n",
            maintype="application", subtype="pdf",
        )
        message.add_attachment(
            b"\x89PNG fictional inline",
            maintype="image", subtype="png", filename="inline.png",
            disposition="inline", cid="<inline-fixture>",
        )
        nested = EmailMessage()
        nested["From"] = "witness@example.invalid"
        nested["To"] = "client@example.invalid"
        nested["Subject"] = "Fictional nested message"
        nested.set_content("Nested fictional evidence.")
        message.add_attachment(nested)
    return message.as_bytes()


class GmailImapTransportTests(unittest.TestCase):
    def adapter(self, raw: bytes):
        client = FakeImapClient({101: (int("1a2b3c", 16), int("4d5e6f", 16), raw)})
        adapter = GmailImapReadonlyAdapter(
            "fictional@example.invalid", "app-password-fixture", lambda _host: client
        )
        self.assertNotIn("app-password-fixture", repr(adapter))
        return adapter, client

    def test_localized_read_only_body_peek_and_supported_attachment(self):
        adapter, client = self.adapter(imap_message())
        result = adapter.intake(
            '{"2032-001"} newer_than:30d -in:spam -in:trash', 5
        )
        self.assertEqual(result["source"], "gmail-imap")
        self.assertTrue(result["coverage"]["complete"])
        self.assertTrue(result["coverage"]["localized_all_mail_discovered"])
        self.assertEqual(result["messages"][0]["id"], "1a2b3c")
        self.assertEqual(result["messages"][0]["thread_id"], "4d5e6f")
        self.assertIn("complete fictional IMAP", result["messages"][0]["body"])
        attachment = result["messages"][0]["attachments"][0]
        self.assertEqual(attachment["filename"], "evidence.txt")
        self.assertEqual(base64.b64decode(attachment["content_base64"]), b"Northstar fictional attachment.\n")
        self.assertEqual(client.selected, ("[Gmail]/Todo el correo", True))
        calls = repr(client.calls)
        self.assertIn("X-GM-RAW", calls)
        self.assertIn("BODY.PEEK[]", calls)
        self.assertNotIn("STORE", calls)
        self.assertEqual(
            {call[0] for call in client.calls},
            {"LOGIN", "LIST", "SELECT", "UID", "LOGOUT"},
        )
        self.assertTrue(
            all(call[1] in {"SEARCH", "FETCH"} for call in client.calls if call[0] == "UID")
        )

    def test_nested_unnamed_and_inline_parts_are_explicit(self):
        adapter, _client = self.adapter(imap_message(include_complex=True))
        result = adapter.intake(
            '{"2032-001"} newer_than:30d -in:spam -in:trash', 5
        )
        attachments = result["messages"][0]["attachments"]
        self.assertNotIn("Nested fictional evidence", result["messages"][0]["body"])
        names = [item["filename"] for item in attachments]
        self.assertIn("attachment-2.pdf", names)
        self.assertIn("inline.png", names)
        self.assertTrue(next(item for item in attachments if item["filename"] == "inline.png")["inline"])
        self.assertTrue(any(name.endswith(".eml") for name in names))
        self.assertTrue(all(item.get("content_base64") for item in attachments))

    def test_oversized_raw_message_is_not_downloaded(self):
        adapter, client = self.adapter(imap_message())
        with patch("lawyeredos.imap_mail.MAX_RAW_MESSAGE_BYTES", 10):
            result = adapter.intake(
                '{"2032-001"} newer_than:30d -in:spam -in:trash', 5
            )
        message = result["messages"][0]
        self.assertFalse(message["body_complete"])
        self.assertIn("bounded", message["attachments"][0]["retrieval_error"])
        self.assertFalse(any("BODY.PEEK[]" in repr(call) for call in client.calls))

    def test_unbounded_query_and_missing_localized_all_mail_fail_closed(self):
        adapter, client = self.adapter(imap_message())
        with self.assertRaisesRegex(ValueError, "bounded matter query"):
            adapter.intake("subject:anything", 5)
        client.list = lambda: ("OK", [b'(\\HasNoChildren) "/" "INBOX"'])
        with self.assertRaisesRegex(ValueError, "localized"):
            adapter.intake('{"2032-001"} newer_than:30d -in:spam -in:trash', 5)

    def test_credentials_are_user_selected_and_no_sidecar_is_read(self):
        config = SimpleNamespace(
            gmail_imap_account_env="FICTIONAL_IMAP_ACCOUNT",
            gmail_imap_secret_env="FICTIONAL_IMAP_SECRET",
            gmail_imap_keychain_service="",
        )
        with patch.dict(
            "os.environ",
            {"FICTIONAL_IMAP_ACCOUNT": "fictional@example.invalid", "FICTIONAL_IMAP_SECRET": "test-secret"},
            clear=True,
        ), patch("lawyeredos.imap_mail.subprocess.run") as keychain:
            self.assertEqual(
                credentials_from_config(config),
                ("fictional@example.invalid", "test-secret"),
            )
        keychain.assert_not_called()

        with patch.dict("os.environ", {}, clear=True):
            with self.assertRaisesRegex(ValueError, "environment variables"):
                credentials_from_config(config)

        config.gmail_imap_keychain_service = "lawyeredos-fictional-test"
        responses = [
            subprocess.CompletedProcess([], 0, '    "acct"<blob>="fictional@example.invalid"\n', ""),
            subprocess.CompletedProcess([], 0, "test-secret\n", ""),
        ]
        with patch.dict("os.environ", {}, clear=True), patch(
            "lawyeredos.imap_mail.subprocess.run", side_effect=responses
        ) as keychain:
            self.assertEqual(
                credentials_from_config(config),
                ("fictional@example.invalid", "test-secret"),
            )
        self.assertEqual(keychain.call_count, 2)
        self.assertNotIn(".gmail-imap-credentials", repr(keychain.call_args_list))

        with patch.dict("os.environ", {}, clear=True), patch(
            "lawyeredos.imap_mail.subprocess.run", side_effect=FileNotFoundError("security")
        ):
            with self.assertRaisesRegex(ValueError, "configured macOS Keychain service"):
                credentials_from_config(config)


if __name__ == "__main__":
    unittest.main()

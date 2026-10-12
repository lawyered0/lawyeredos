"""Fictional fixtures for the remaining recovered document claims."""
import importlib.util
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]


def script(name):
    sys.path.insert(0, str(ROOT / "scripts"))
    try:
        spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / (name + ".py"))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.pop(0)


class RecoveredDocumentRegressionTests(unittest.TestCase):
    def test_field_values_cannot_inject_another_placeholder(self):
        creator = script("create_letter")
        document = Document()
        paragraph = document.add_paragraph("ORGANIZATION NAME {{DATE}}")
        with self.assertRaisesRegex(ValueError, "field values"):
            creator.replace_simple(document, {"ORGANIZATION NAME": "{{DATE}}", "{{DATE}}": "Fictional date"})
        self.assertEqual(paragraph.text, "ORGANIZATION NAME {{DATE}}")

    def test_placeholder_value_containing_its_key_terminates(self):
        code = '''from docx import Document
from create_letter import _replace_in_runs
document = Document()
paragraph = document.add_paragraph("ORGANIZATION NAME and ORGANIZATION NAME")
_replace_in_runs(paragraph, "ORGANIZATION NAME", "NORTH ORGANIZATION NAME SERVICES")
assert paragraph.text == "NORTH ORGANIZATION NAME SERVICES and NORTH ORGANIZATION NAME SERVICES"
'''
        completed = subprocess.run([sys.executable, "-c", code], cwd=ROOT / "scripts", capture_output=True, text=True, timeout=3)
        self.assertEqual(completed.returncode, 0, completed.stderr)

    def test_acceptance_refuses_structural_deletions_without_output(self):
        acceptor = script("accept_tracked_changes")
        for location, marker in (("paragraph", "del"), ("paragraph", "moveFrom"), ("row", "del"), ("row", "moveFrom")):
            with self.subTest(location=location, marker=marker), tempfile.TemporaryDirectory() as temporary:
                folder = Path(temporary)
                source, output = folder / "fictional.docx", folder / "clean.docx"
                document = Document()
                if location == "paragraph":
                    paragraph = document.add_paragraph("Merge this fictional paragraph.")
                    document.add_paragraph("With the following paragraph.")
                    props = paragraph._p.get_or_add_pPr()
                    run_props = OxmlElement("w:rPr")
                    run_props.append(OxmlElement("w:" + marker))
                    props.append(run_props)
                else:
                    table = document.add_table(rows=2, cols=1)
                    table.cell(0, 0).text = "Kept fictional row."
                    table.cell(1, 0).text = "Deleted fictional row."
                    props = table.rows[1]._tr.get_or_add_trPr()
                    props.append(OxmlElement("w:" + marker))
                document.save(source)
                before = source.read_bytes()
                with self.assertRaisesRegex(ValueError, "structural.*deletion"):
                    acceptor.accept(source, output)
                self.assertFalse(output.exists())
                self.assertEqual(source.read_bytes(), before)

    def test_scrubber_refuses_review_metadata_without_erasing_content(self):
        scrubber = script("scrub_docx_metadata")
        for kind in ("comments", "people", "revision"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temporary:
                folder = Path(temporary)
                source, output = folder / "fictional.docx", folder / "scrubbed.docx"
                document = Document()
                paragraph = document.add_paragraph("Fictional clause.")
                if kind == "revision":
                    insertion = OxmlElement("w:ins")
                    insertion.set(qn("w:author"), "Fictional Reviewer")
                    paragraph._p.append(insertion)
                document.save(source)
                if kind != "revision":
                    with zipfile.ZipFile(source, "a") as package:
                        package.writestr("word/" + kind + ".xml", '<records author="Fictional Reviewer"/>')
                before = source.read_bytes()
                with self.assertRaisesRegex(ValueError, "review metadata"):
                    scrubber.scrub(source, output)
                self.assertFalse(output.exists())
                self.assertEqual(source.read_bytes(), before)

    def test_scrubber_clears_additional_core_metadata(self):
        scrubber = script("scrub_docx_metadata")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source, output = folder / "fictional.docx", folder / "scrubbed.docx"
            document = Document()
            document.add_paragraph("Fictional clean correspondence.")
            props = document.core_properties
            for field in ("title", "subject", "keywords", "comments", "category", "identifier", "content_status"):
                setattr(props, field, "Fictional confidential metadata")
            props.revision = 47
            document.save(source)
            scrubber.scrub(source, output)
            with zipfile.ZipFile(output) as package:
                core = etree.fromstring(package.read("docProps/core.xml"))
                self.assertNotIn(b"Fictional confidential metadata", etree.tostring(core))
                self.assertFalse(any(node.tag.rsplit("}", 1)[-1] == "revision" for node in core))

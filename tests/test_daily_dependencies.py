from __future__ import annotations

import argparse
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from lawyeredos.common import sha256_file
from lawyeredos.integrations import GmailReadonlyAdapter, GoogleCalendarAdapter
from lawyeredos.memory_search import build_index, memory_commit, memory_preview, search
from lawyeredos.review_bundle import create_manifest
from lawyeredos.review_orchestrator import LENSES, collect, prepare
from lawyeredos.tracker_cli import initialize
from lawyeredos.tracker_migration import commit_migration, export_tracker, migration_plan
from lawyeredos.tracker_xlsx import project
from lawyeredos.voice_lint import lint


class DailyDependencyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def source_json(self) -> Path:
        path = self.root / "legacy.json"
        path.write_text(json.dumps({"matters": [{
            "File #": "2031-001", "Client Name": "Example Orchard Ltd.",
            "Matter Description": "Fictional equipment dispute",
            "Matter Folder": "2031-001-example-orchard", "Status": "Open Matters",
            "Court Deadlines": "legacy note for review",
        }]}), encoding="utf-8")
        return path

    def migrate(self) -> Path:
        source = self.source_json(); target = self.root / "matters.sqlite3"
        plan = migration_plan(source, target)
        result = commit_migration(source, target, plan["confirm_token"])
        self.assertTrue(result["committed"])
        return target

    def test_safe_migration_exports_and_xlsx_projection(self):
        database = self.migrate()
        with sqlite3.connect(database) as connection:
            count = connection.execute("SELECT COUNT(*) FROM matters").fetchone()[0]
        self.assertEqual(count, 1)
        export = self.root / "tracker.json"
        self.assertEqual(export_tracker(database, export, "json")["matter_count"], 1)
        xlsx = self.root / "tracker.xlsx"
        projected = project(database, xlsx)
        self.assertEqual(projected["counts"]["open"], 1)
        from openpyxl import load_workbook
        workbook = load_workbook(xlsx, read_only=True)
        self.assertEqual(workbook.sheetnames, ["Open Matters", "Closed Matters"])
        self.assertEqual(workbook["Open Matters"]["A2"].value, "2031-001")

    def test_migration_never_overwrites(self):
        target = self.root / "exists.sqlite3"; target.write_text("leave me", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "already exists"):
            migration_plan(self.source_json(), target)
        self.assertEqual(target.read_text(encoding="utf-8"), "leave me")

    def test_gmail_read_only_adapter_uses_fake_transport(self):
        calls = []
        def transport(method, url, headers, body):
            calls.append((method, url, headers, body))
            if "/messages?" in url:
                return 200, json.dumps({"messages": [{"id": "m1"}]}).encode()
            return 200, json.dumps({"id": "m1", "threadId": "t1", "snippet": "fictional",
                "payload": {"headers": [{"name": "Subject", "value": "Example"}]}}).encode()
        result = GmailReadonlyAdapter("test-token", transport).search("after:2031/01/01", 5)
        self.assertTrue(result["read_only"])
        self.assertEqual(result["messages"][0]["subject"], "Example")
        self.assertTrue(all(call[0] == "GET" and call[3] is None for call in calls))

    def test_calendar_preview_and_commit_use_fake_transport(self):
        calls = []
        def transport(method, url, headers, body):
            calls.append((method, url, body))
            if method == "GET" and "events?" in url: return 200, b'{"items":[]}'
            if method == "GET": return 200, b'{"id":"event-1","extendedProperties":{"private":{"sync_key":"2031-001:COURT:fictional"}}}'
            return 200, b'{"id":"event-1","status":"confirmed"}'
        adapter = GoogleCalendarAdapter("test-token", "deadlines@example.invalid", transport)
        listing = adapter.list_events("2031-01-01T00:00:00Z", "2031-02-01T00:00:00Z")
        self.assertEqual(listing["coverage"], "complete")
        event = {"summary": "COURT | 2031-001 | fictional", "extendedProperties": {"private": {"sync_key": "2031-001:COURT:fictional"}}}
        preview = adapter.preview("create", event)
        committed = adapter.commit("create", event, "", preview["confirm_token"])
        self.assertTrue(committed["committed"])
        self.assertEqual([call[0] for call in calls], ["GET", "POST", "GET"])
        with self.assertRaisesRegex(ValueError, "dedicated"):
            GoogleCalendarAdapter("test-token", "primary", transport).list_events("a", "b")

    def test_memory_search_and_compare_and_swap(self):
        matter = self.root / "matter"; matter.mkdir()
        brief = matter / "_matter-brief.md"
        brief.write_text("# Example\nThe delivery date is disputed.\n", encoding="utf-8")
        index = self.root / "search.sqlite3"
        build_index(matter, index)
        self.assertEqual(search(index, "delivery")["results"][0]["path"], "_matter-brief.md")
        content = self.root / "new.md"; content.write_text("# Example\nUpdated fictional note.\n", encoding="utf-8")
        digest = sha256_file(brief)
        preview = memory_preview(self.root, "matter/_matter-brief.md", content.read_bytes(), digest)
        result = memory_commit(self.root, "matter/_matter-brief.md", content.read_bytes(), digest, preview["confirm_token"])
        self.assertTrue(result["committed"])
        with self.assertRaisesRegex(ValueError, "compare-and-swap"):
            memory_preview(self.root, "matter/_matter-brief.md", b"later", digest)

    def test_voice_lint(self):
        result = lint("Please pay 5,000 in settlement. This is a short sentence.", "email")
        self.assertFalse(result["ok"])
        self.assertEqual(result["findings"][0]["rule"], "currency")

    def test_review_orchestration_with_synthetic_results(self):
        source = self.root / "source.txt"; source.write_text("Fictional source.", encoding="utf-8")
        draft = self.root / "draft.txt"; draft.write_text("Fictional draft.", encoding="utf-8")
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(create_manifest(draft, [source], self.root)), encoding="utf-8")
        tasks = self.root / "tasks.json"
        self.assertEqual(prepare(manifest, tasks)["task_count"], 4)
        task_bundle = json.loads(tasks.read_text(encoding="utf-8"))
        for task in task_bundle["tasks"]:
            self.assertEqual(task["result_schema"]["severity_allowed"],
                             ["critical", "material", "optional"])
        results = self.root / "results"; results.mkdir()
        for lens in LENSES:
            (results / f"{lens}.json").write_text(json.dumps({
                "lens": lens, "coverage_gaps": [], "findings": [{
                    "severity": "material", "title": "Synthetic issue",
                    "observation": "The draft differs from the source.",
                    "source_path": str(source.resolve()), "location": "line 1",
                    "countervailing_evidence": "The supplied draft was checked.",
                    "consequence": "A reader could misunderstand the fact.",
                }]}), encoding="utf-8")
        report = self.root / "report.json"
        result = collect(tasks, results, report)
        self.assertFalse(result["ok"])
        self.assertTrue(result["collection_complete"])
        self.assertFalse(result["ready"])
        self.assertFalse(result["authorizes_action"])
        self.assertIn("critical or material findings remain unresolved", result["readiness_reasons"])
        self.assertTrue(result["deliverable_unchanged"])
        self.assertEqual(result["finding_count"], 4)
        invalid_path = results / f"{next(iter(LENSES))}.json"
        invalid = json.loads(invalid_path.read_text(encoding="utf-8"))
        invalid["findings"][0]["severity"] = "minor"
        invalid_path.write_text(json.dumps(invalid), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "invalid severity; expected one of"):
            collect(tasks, results, self.root / "invalid-report.json")

    def test_review_readiness_blocks_missing_sources_gaps_and_questions(self):
        draft = self.root / "draft.txt"
        draft.write_text("Fictional draft.", encoding="utf-8")
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(create_manifest(draft, [], self.root)), encoding="utf-8")
        tasks = self.root / "tasks.json"
        prepare(manifest, tasks)
        results = self.root / "results"
        results.mkdir()
        for index, lens in enumerate(LENSES):
            (results / f"{lens}.json").write_text(json.dumps({
                "lens": lens,
                "coverage_gaps": ["fictional attachment unavailable"] if index == 0 else [],
                "questions": ["Which fictional date controls?"] if index == 1 else [],
                "findings": [],
            }), encoding="utf-8")

        collected = collect(tasks, results, self.root / "report.json")
        self.assertFalse(collected["ok"])
        self.assertTrue(collected["collection_complete"])
        self.assertFalse(collected["ready"])
        self.assertEqual(collected["blocking_finding_count"], 0)
        self.assertEqual(collected["readiness_reasons"], [
            "no controlling sources were supplied",
            "review coverage gaps remain",
            "review questions remain unresolved",
        ])

    def test_review_with_sources_and_only_optional_findings_can_be_ready_but_never_authorizes(self):
        source = self.root / "source.txt"
        source.write_text("Fictional source.", encoding="utf-8")
        draft = self.root / "draft.txt"
        draft.write_text("Fictional draft.", encoding="utf-8")
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(create_manifest(draft, [source], self.root)), encoding="utf-8")
        tasks = self.root / "tasks.json"
        prepare(manifest, tasks)
        results = self.root / "results"
        results.mkdir()
        for lens in LENSES:
            (results / f"{lens}.json").write_text(json.dumps({
                "lens": lens,
                "coverage_gaps": [],
                "questions": [],
                "findings": [{
                    "severity": "optional",
                    "title": "Style refinement",
                    "observation": "A fictional sentence could be shorter.",
                    "source_path": str(draft.resolve()),
                    "location": "line 1",
                    "countervailing_evidence": "The current sentence remains understandable.",
                    "consequence": "This affects style only.",
                }],
            }), encoding="utf-8")

        collected = collect(tasks, results, self.root / "report.json")
        self.assertTrue(collected["collection_complete"])
        self.assertTrue(collected["ready"])
        self.assertFalse(collected["authorizes_action"])

    def test_review_rejects_empty_or_subset_task_bundles(self):
        source = self.root / "source.txt"
        source.write_text("Fictional source.", encoding="utf-8")
        draft = self.root / "draft.txt"
        draft.write_text("Fictional draft.", encoding="utf-8")
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(create_manifest(draft, [source], self.root)), encoding="utf-8")
        original_tasks = self.root / "original-tasks.json"
        prepare(manifest, original_tasks)
        bundle = json.loads(original_tasks.read_text(encoding="utf-8"))
        results = self.root / "results"
        results.mkdir()

        for name, tasks in (("empty", []), ("subset", bundle["tasks"][:-1])):
            tampered = {**bundle, "tasks": tasks}
            tasks_path = self.root / f"{name}-tasks.json"
            tasks_path.write_text(json.dumps(tampered), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "each expected lens exactly once"):
                collect(tasks_path, results, self.root / f"{name}-report.json")

    def test_review_rejects_schema_not_bound_to_host(self):
        source = self.root / "source.txt"
        source.write_text("Fictional source.", encoding="utf-8")
        draft = self.root / "draft.txt"
        draft.write_text("Fictional draft.", encoding="utf-8")
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(create_manifest(draft, [source], self.root)), encoding="utf-8")
        tasks = self.root / "tasks.json"
        prepare(manifest, tasks, host="claude")
        bundle = json.loads(tasks.read_text(encoding="utf-8"))
        bundle["host"] = "codex"
        tasks.write_text(json.dumps(bundle), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "invalid schema"):
            collect(tasks, self.root / "results", self.root / "report.json")

    def test_review_rejects_tampered_task_authority_and_evidence(self):
        source = self.root / "source.txt"
        source.write_text("Fictional source.", encoding="utf-8")
        other = self.root / "other.txt"
        other.write_text("Different fictional evidence.", encoding="utf-8")
        draft = self.root / "draft.txt"
        draft.write_text("Fictional draft.", encoding="utf-8")
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(create_manifest(draft, [source], self.root)), encoding="utf-8")
        original_tasks = self.root / "tasks.json"
        prepare(manifest, original_tasks)
        original = json.loads(original_tasks.read_text(encoding="utf-8"))

        mutations = {
            "agent": lambda task: task.__setitem__("agent", "untrusted-agent"),
            "read-only": lambda task: task.__setitem__("read_only", False),
            "evidence": lambda task: task.__setitem__(
                "sources", [{"path": str(other.resolve()), "size": other.stat().st_size, "sha256": "0" * 64}]
            ),
            "result-schema": lambda task: task["result_schema"].__setitem__(
                "severity_allowed", ["optional"]
            ),
        }
        for name, mutate in mutations.items():
            with self.subTest(name=name):
                bundle = json.loads(json.dumps(original))
                mutate(bundle["tasks"][0])
                tampered = self.root / f"tampered-{name}.json"
                tampered.write_text(json.dumps(bundle), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "does not match the canonical"):
                    collect(tampered, self.root / "results", self.root / f"report-{name}.json")

    def test_review_rejects_cross_host_manifest_at_prepare_and_collect(self):
        source = self.root / "source.txt"
        source.write_text("Fictional source.", encoding="utf-8")
        draft = self.root / "draft.txt"
        draft.write_text("Fictional draft.", encoding="utf-8")
        claude_manifest = self.root / "claude-manifest.json"
        claude_manifest.write_text(
            json.dumps(create_manifest(draft, [source], self.root, host="claude")),
            encoding="utf-8",
        )
        with self.assertRaisesRegex(ValueError, "manifest schema does not match"):
            prepare(claude_manifest, self.root / "codex-tasks.json", host="codex")

        tasks = self.root / "claude-tasks.json"
        prepare(claude_manifest, tasks, host="claude")
        bundle = json.loads(tasks.read_text(encoding="utf-8"))
        bundle["host"] = "codex"
        bundle["schema"] = "codex-review-tasks-v1"
        for task in bundle["tasks"]:
            task["agent"] = "independent-codex-reviewer"
        tasks.write_text(json.dumps(bundle), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "manifest schema does not match"):
            collect(tasks, self.root / "results", self.root / "report.json")


if __name__ == "__main__":
    unittest.main()

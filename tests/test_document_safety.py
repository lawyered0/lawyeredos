from __future__ import annotations

import base64
import hashlib
import importlib.util
import io
import json
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock
from xml.etree import ElementTree


ROOT = Path(__file__).resolve().parents[1]


def load_script(name: str):
    path = ROOT / "scripts" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(f"document_safety_{name}", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def letter_args(template: Path, body: Path, output: Path):
    manifest = template.parent / "template-manifest.json"
    manifest.write_text(json.dumps({
        "schema_version": 1,
        "algorithm": "sha256",
        "template": template.name,
        "sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
    }), encoding="utf-8")
    return type("Args", (), {
        "template": str(template),
        "template_manifest": str(manifest),
        "body_file": str(body),
        "output": str(output),
        "organization": "Northstar Legal Lab",
        "date": "October 11, 2026",
        "recipient": "Alex Example",
        "address": "100 Example Street",
        "subject": "Fictional matter",
        "signoff": "Example Counsel",
        "author": "Authorized reviewer",
    })()


class DocumentSafetyTests(unittest.TestCase):
    def setUp(self):
        try:
            from docx import Document  # noqa: F401
            from lxml import etree  # noqa: F401
        except ImportError:
            self.skipTest("document dependencies are not installed")

    def _template(self, folder: Path) -> Path:
        from docx import Document

        path = folder / "template.docx"
        document = Document()
        date = document.add_paragraph()
        lead = date.add_run("Date: ")
        lead.italic = True
        marker = date.add_run("{{DA")
        marker.bold = True
        date.add_run("TE}}")
        tail = date.add_run(" — private draft")
        tail.underline = True
        document.add_paragraph("{{BODY}}")
        header_cell = document.sections[0].header.add_table(1, 1, 6_000_000).cell(0, 0)
        header = header_cell.paragraphs[0]
        header.add_run("ORGAN")
        header.add_run("IZATION NAME")
        document.save(path)
        return path

    def test_letter_replaces_split_markers_without_flattening_styles(self):
        from docx import Document

        creator = load_script("create_letter")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            template = self._template(folder)
            body = folder / "body.txt"
            body.write_text("This is fictional correspondence.", encoding="utf-8")
            output = folder / "letter.docx"
            creator.create(letter_args(template, body, output))

            document = Document(output)
            date = document.paragraphs[0]
            self.assertEqual(date.text, "Date: October 11, 2026 — private draft")
            self.assertTrue(date.runs[0].italic)
            self.assertTrue(date.runs[1].bold)
            self.assertTrue(date.runs[-1].underline)
            header = document.sections[0].header.tables[0].cell(0, 0).text
            self.assertEqual(header, "Northstar Legal Lab")

    def test_letter_refuses_review_markup_and_does_not_publish(self):
        from docx import Document
        from docx.oxml import OxmlElement

        creator = load_script("create_letter")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            template = self._template(folder)
            document = Document(template)
            document.paragraphs[0]._p.append(OxmlElement("w:ins"))
            document.save(template)
            body = folder / "body.txt"
            body.write_text("Fictional body.", encoding="utf-8")
            output = folder / "letter.docx"
            with self.assertRaisesRegex(ValueError, "pre-existing revision/comment"):
                creator.create(letter_args(template, body, output))
            self.assertFalse(output.exists())

            linked_template = folder / "linked-template.docx"
            linked_template.symlink_to(template)
            with self.assertRaisesRegex(ValueError, "template must be a regular, non-symlink file"):
                creator.create(letter_args(linked_template, body, output))
            self.assertFalse(output.exists())

            property_template = self._template(folder)
            document = Document(property_template)
            paragraph = document.paragraphs[0]
            paragraph._p.get_or_add_pPr().append(OxmlElement("w:tblPrExChange"))
            document.save(property_template)
            with self.assertRaisesRegex(ValueError, "revision marker tblPrExChange"):
                creator.create(letter_args(property_template, body, output))
            self.assertFalse(output.exists())

    def test_acceptance_removes_revisions_and_track_revisions_setting(self):
        from docx import Document

        vendor = ROOT / "vendor" / "legal-redline-tools"
        import sys
        sys.path.insert(0, str(vendor))
        try:
            from legal_redline import apply_redlines
        finally:
            sys.path.pop(0)
        accepter = load_script("accept_tracked_changes")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "source.docx"
            redlined = folder / "redlined.docx"
            accepted = folder / "accepted.docx"
            document = Document()
            document.add_paragraph("Payment is due in thirty days.")
            document.save(source)
            apply_redlines(str(source), str(redlined), [{
                "type": "replace", "old": "thirty days", "new": "forty-five days",
            }])
            accepter.accept(redlined, accepted)
            local_tags = set()
            with zipfile.ZipFile(accepted) as package:
                for name in package.namelist():
                    if name.startswith("word/") and name.endswith(".xml"):
                        root = ElementTree.fromstring(package.read(name))
                        local_tags.update(element.tag.rsplit("}", 1)[-1] for element in root.iter())
            for marker in ("ins", "del", "trackRevisions"):
                self.assertNotIn(marker, local_tags)
            self.assertIn("forty-five days", Document(accepted).paragraphs[0].text)

    def test_acceptance_refuses_comments_and_unsupported_property_revisions(self):
        from docx import Document
        from docx.oxml import OxmlElement

        accepter = load_script("accept_tracked_changes")
        for marker in ("w:commentReference", "w:pPrChange", "w:tblPrExChange"):
            with self.subTest(marker=marker), tempfile.TemporaryDirectory() as temporary:
                folder = Path(temporary)
                source = folder / "source.docx"
                output = folder / "accepted.docx"
                document = Document()
                paragraph = document.add_paragraph("Fictional clause.")
                if marker == "w:pPrChange":
                    paragraph._p.get_or_add_pPr().append(OxmlElement(marker))
                else:
                    paragraph._p.append(OxmlElement(marker))
                document.save(source)
                with self.assertRaisesRegex(ValueError, "cannot be safely accepted"):
                    accepter.accept(source, output)
                self.assertFalse(output.exists())

    def test_acceptance_refuses_deleted_embedded_content(self):
        from docx import Document
        from docx.oxml import OxmlElement

        accepter = load_script("accept_tracked_changes")
        pixel = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "source.docx"
            output = folder / "accepted.docx"
            document = Document()
            paragraph = document.add_paragraph("Deleted exhibit: ")
            run = paragraph.add_run()
            run.add_picture(io.BytesIO(pixel))
            paragraph._p.remove(run._r)
            deletion = OxmlElement("w:del")
            deletion.append(run._r)
            paragraph._p.append(deletion)
            document.save(source)
            with zipfile.ZipFile(source) as package:
                self.assertTrue(any(name.startswith("word/media/") for name in package.namelist()))
            with self.assertRaisesRegex(ValueError, "external/package content"):
                accepter.accept(source, output)
            self.assertFalse(output.exists())

    def test_letter_output_collision_and_symlink_leave_target_unchanged(self):
        creator = load_script("create_letter")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            template = self._template(folder)
            body = folder / "body.txt"
            body.write_text("Fictional body.", encoding="utf-8")
            existing = folder / "existing.docx"
            existing.write_bytes(b"do not replace")
            with self.assertRaisesRegex(ValueError, "output already exists"):
                creator.create(letter_args(template, body, existing))
            self.assertEqual(existing.read_bytes(), b"do not replace")

            victim = folder / "victim.docx"
            victim.write_bytes(b"victim")
            link = folder / "linked-output.docx"
            link.symlink_to(victim)
            with self.assertRaisesRegex(ValueError, "output already exists"):
                creator.create(letter_args(template, body, link))
            self.assertEqual(victim.read_bytes(), b"victim")

            target_dir = folder / "target" / "nested"
            target_dir.mkdir(parents=True)
            alias = folder / "alias"
            alias.symlink_to(folder / "target", target_is_directory=True)
            nested_output = alias / "nested" / "letter.docx"
            with self.assertRaisesRegex(ValueError, "must not contain a symlink"):
                creator.create(letter_args(template, body, nested_output))
            self.assertFalse((target_dir / "letter.docx").exists())

            missing_output = alias / "new-directory" / "letter.docx"
            with self.assertRaisesRegex(ValueError, "must not contain a symlink"):
                creator.create(letter_args(template, body, missing_output))
            self.assertFalse((folder / "target" / "new-directory").exists())

    def test_acceptance_output_collision_and_symlink_leave_target_unchanged(self):
        from docx import Document

        accepter = load_script("accept_tracked_changes")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "source.docx"
            document = Document()
            document.add_paragraph("Fictional clean clause.")
            document.save(source)

            existing = folder / "existing.docx"
            existing.write_bytes(b"do not replace")
            with self.assertRaisesRegex(ValueError, "output already exists"):
                accepter.accept(source, existing)
            self.assertEqual(existing.read_bytes(), b"do not replace")

            victim = folder / "victim.docx"
            victim.write_bytes(b"victim")
            link = folder / "linked-output.docx"
            link.symlink_to(victim)
            with self.assertRaisesRegex(ValueError, "output already exists"):
                accepter.accept(source, link)
            self.assertEqual(victim.read_bytes(), b"victim")

            target_dir = folder / "target" / "nested"
            target_dir.mkdir(parents=True)
            alias = folder / "alias"
            alias.symlink_to(folder / "target", target_is_directory=True)
            nested_output = alias / "nested" / "accepted.docx"
            with self.assertRaisesRegex(ValueError, "must not contain a symlink"):
                accepter.accept(source, nested_output)
            self.assertFalse((target_dir / "accepted.docx").exists())

            missing_output = alias / "new-directory" / "accepted.docx"
            with self.assertRaisesRegex(ValueError, "must not contain a symlink"):
                accepter.accept(source, missing_output)
            self.assertFalse((folder / "target" / "new-directory").exists())

    def test_script_cleanup_preserves_concurrent_replacement(self):
        for helper_name in ("create_letter", "accept_tracked_changes"):
            with self.subTest(helper=helper_name), tempfile.TemporaryDirectory() as temporary:
                helper = load_script(helper_name)
                output = Path(temporary) / "output.docx"

                def replace_then_fail(_descriptor):
                    output.write_bytes(b"concurrent replacement")
                    raise OSError("fictional fsync failure")

                with mock.patch.object(helper.os, "fsync", side_effect=replace_then_fail):
                    with self.assertRaisesRegex(OSError, "fictional fsync failure"):
                        helper._write_new(output, b"generated bytes")
                self.assertEqual(output.read_bytes(), b"concurrent replacement")

    def test_letter_scans_whole_package_drafting_markers_but_allows_citations(self):
        from docx import Document

        creator = load_script("create_letter")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            template = self._template(folder)
            document = Document(template)
            document.sections[0].footer.paragraphs[0].text = "[VERIFY]"
            document.save(template)
            body = folder / "body.txt"
            body.write_text("See Example v Sample, [2024] 1 Fictional R 2.", encoding="utf-8")
            output = folder / "letter.docx"
            with self.assertRaisesRegex(ValueError, r"\[VERIFY\]"):
                creator.create(letter_args(template, body, output))
            self.assertFalse(output.exists())

            template = self._template(folder)
            body.write_text("Confirm the TK deadline.", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "drafting marker: TK"):
                creator.create(letter_args(template, body, output))
            self.assertFalse(output.exists())

            body.write_text("See Example v Sample, [2024] 1 Fictional R 2.", encoding="utf-8")
            creator.create(letter_args(template, body, output))
            self.assertTrue(output.is_file())

    def test_letter_uses_the_same_template_bytes_it_scanned(self):
        from docx import Document

        creator = load_script("create_letter")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            template = self._template(folder)
            body = folder / "body.txt"
            body.write_text("Fictional body.", encoding="utf-8")
            output = folder / "letter.docx"
            original_scan = creator.package_review_markers
            calls = 0

            def swap_after_scan(package):
                nonlocal calls
                result = original_scan(package)
                calls += 1
                if calls == 1:
                    replacement = Document()
                    replacement.add_paragraph("SWAPPED SOURCE")
                    replacement.add_paragraph("{{BODY}}")
                    replacement.save(template)
                return result

            with mock.patch.object(creator, "package_review_markers", side_effect=swap_after_scan):
                creator.create(letter_args(template, body, output))
            rendered = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
            self.assertNotIn("SWAPPED SOURCE", rendered)
            self.assertIn("Fictional body.", rendered)

    def test_letter_refuses_template_hash_or_path_not_bound_by_manifest(self):
        creator = load_script("create_letter")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            template = self._template(folder)
            body = folder / "body.txt"
            body.write_text("Fictional body.", encoding="utf-8")
            args = letter_args(template, body, folder / "hash-failure.docx")
            manifest = Path(args.template_manifest)
            value = json.loads(manifest.read_text(encoding="utf-8"))
            value["sha256"] = "0" * 64
            manifest.write_text(json.dumps(value), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "template integrity check failed"):
                creator.create(args)
            self.assertFalse(Path(args.output).exists())

            args = letter_args(template, body, folder / "path-failure.docx")
            value = json.loads(manifest.read_text(encoding="utf-8"))
            value["template"] = "different.docx"
            manifest.write_text(json.dumps(value), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "path does not match"):
                creator.create(args)
            self.assertFalse(Path(args.output).exists())

    def test_document_clis_sanitize_truncated_packages_without_output(self):
        creator = load_script("create_letter")
        accepter = load_script("accept_tracked_changes")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            broken = folder / "broken.docx"
            broken.write_bytes(b"PK\x03\x04truncated")
            body = folder / "body.txt"
            body.write_text("Fictional body.", encoding="utf-8")
            letter_output = folder / "letter.docx"
            clean_output = folder / "clean.docx"
            manifest = folder / "template-manifest.json"
            manifest.write_text(json.dumps({
                "schema_version": 1, "algorithm": "sha256", "template": broken.name,
                "sha256": hashlib.sha256(broken.read_bytes()).hexdigest(),
            }), encoding="utf-8")

            creator_argv = [
                "create_letter.py", "--template", str(broken),
                "--template-manifest", str(manifest), "--body-file", str(body),
                "--output", str(letter_output), "--organization", "Northstar Legal Lab",
                "--date", "October 11, 2026", "--recipient", "Alex Example",
                "--address", "100 Example Street", "--subject", "Fictional matter",
                "--signoff", "Example Counsel",
            ]
            output = io.StringIO()
            with mock.patch("sys.argv", creator_argv), redirect_stdout(output):
                self.assertEqual(creator.main(), 2)
            self.assertIn("error:", output.getvalue())
            self.assertFalse(letter_output.exists())

            output = io.StringIO()
            with mock.patch("sys.argv", [
                "accept_tracked_changes.py", str(broken), "--output", str(clean_output),
            ]), redirect_stdout(output):
                self.assertEqual(accepter.main(), 2)
            self.assertIn("error:", output.getvalue())
            self.assertFalse(clean_output.exists())

            with zipfile.ZipFile(broken, "w") as package:
                package.writestr("notes.txt", "not a Word package")
            value = json.loads(manifest.read_text(encoding="utf-8"))
            value["sha256"] = hashlib.sha256(broken.read_bytes()).hexdigest()
            manifest.write_text(json.dumps(value), encoding="utf-8")
            output = io.StringIO()
            with mock.patch("sys.argv", creator_argv), redirect_stdout(output):
                self.assertEqual(creator.main(), 2)
            self.assertIn("not a valid Word DOCX package", output.getvalue())
            self.assertFalse(letter_output.exists())
            output = io.StringIO()
            with mock.patch("sys.argv", [
                "accept_tracked_changes.py", str(broken), "--output", str(clean_output),
            ]), redirect_stdout(output):
                self.assertEqual(accepter.main(), 2)
            self.assertIn("not a valid Word DOCX package", output.getvalue())
            self.assertFalse(clean_output.exists())

            with zipfile.ZipFile(broken, "w") as package:
                package.writestr("[Content_Types].xml", "<Types/>")
                package.writestr("_rels/.rels", "<Relationships/>")
                package.writestr("word/document.xml", "<document/>")
            value["sha256"] = hashlib.sha256(broken.read_bytes()).hexdigest()
            manifest.write_text(json.dumps(value), encoding="utf-8")
            output = io.StringIO()
            with mock.patch("sys.argv", creator_argv), redirect_stdout(output):
                self.assertEqual(creator.main(), 2)
            self.assertIn("not a valid Word DOCX package", output.getvalue())
            self.assertFalse(letter_output.exists())
            output = io.StringIO()
            with mock.patch("sys.argv", [
                "accept_tracked_changes.py", str(broken), "--output", str(clean_output),
            ]), redirect_stdout(output):
                self.assertEqual(accepter.main(), 2)
            self.assertIn("not a valid Word DOCX package", output.getvalue())
            self.assertFalse(clean_output.exists())

    def test_document_snapshot_limits_fail_before_output(self):
        creator = load_script("create_letter")
        accepter = load_script("accept_tracked_changes")
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            oversized = folder / "oversized"
            oversized.write_bytes(b"x" * 65)
            with mock.patch.object(creator, "MAX_DOCX_BYTES", 64):
                with self.assertRaisesRegex(ValueError, "64-byte safety limit"):
                    creator._read_regular_bytes(oversized, "template", creator.MAX_DOCX_BYTES)
            with mock.patch.object(accepter, "MAX_DOCX_BYTES", 64):
                with self.assertRaisesRegex(ValueError, "64-byte safety limit"):
                    accepter._read_regular_bytes(oversized)

            template = self._template(folder)
            with zipfile.ZipFile(template) as package:
                with mock.patch.object(creator, "MAX_DOCX_UNCOMPRESSED_BYTES", 64):
                    with self.assertRaisesRegex(ValueError, "uncompressed safety limit"):
                        creator._validate_docx_package(package)


if __name__ == "__main__":
    unittest.main()

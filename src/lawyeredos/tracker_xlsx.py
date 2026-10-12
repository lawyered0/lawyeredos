from __future__ import annotations

import argparse
import io
import json
import sqlite3
from pathlib import Path
from typing import Any

from .common import emit, sqlite_ro_uri
from .private_fs import create_private_file
from .tracker_values import ILLEGAL_XML_CONTROLS


HEADERS = [
    ("file_number", "File #"), ("client", "Client"),
    ("description", "Matter Description"), ("status", "Status"),
    ("next_action", "Next Action"), ("next_action_date", "Next Action Date"),
    ("limitation_date", "Limitation Date"), ("court_deadlines", "Court Deadlines"),
    ("opposing_party", "Opposing Party"), ("other_parties", "Other Parties"),
    ("primary_email", "Primary Email"), ("matter_folder", "Matter Folder"),
    ("updated_at", "Updated At"),
]


def _rows(database: Path, status: str) -> list[dict[str, Any]]:
    uri = sqlite_ro_uri(database)
    with sqlite3.connect(uri, uri=True) as connection:
        connection.row_factory = sqlite3.Row
        return [dict(row) for row in connection.execute(
            "SELECT * FROM matters WHERE status=? ORDER BY file_number", (status,)
        )]


def project(database: Path, output: Path) -> dict[str, Any]:
    if output.exists():
        raise ValueError("XLSX output already exists")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise ValueError("openpyxl is required for XLSX projection") from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts: dict[str, int] = {}
    sanitized: list[dict[str, Any]] = []
    truncated: list[dict[str, Any]] = []
    omitted_deadline_fields: list[dict[str, Any]] = []
    for status, title in (("open", "Open Matters"), ("closed", "Closed Matters")):
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.append([label for _, label in HEADERS])
        rows = _rows(database, status)
        counts[status] = len(rows)
        for row in rows:
            values = []
            for key, _label in HEADERS:
                value = row.get(key, "")
                if key == "court_deadlines":
                    parsed = json.loads(value or "[]")
                    for index, item in enumerate(parsed):
                        if isinstance(item, dict):
                            omitted = sorted(set(item) - {"date", "description"})
                            if omitted:
                                omitted_deadline_fields.append({"file_number": row["file_number"],
                                    "field": key, "item_index": index, "omitted_fields": omitted})
                    value = "\n".join(
                        " | ".join(str(item.get(name, "")) for name in ("date", "description") if item.get(name))
                        if isinstance(item, dict) else str(item) for item in parsed
                    )
                if isinstance(value, str):
                    changed, removed = ILLEGAL_XML_CONTROLS.subn("\ufffd", value)
                    if removed:
                        sanitized.append({"file_number": row["file_number"], "field": key, "removed_characters": removed})
                    value = changed
                    if len(value) > 32767:
                        truncated.append({"file_number": row["file_number"], "field": key, "original_characters": len(value)})
                        marker = "\n[truncated; full text in tracker DB]"
                        value = value[:32767 - len(marker)] + marker
                values.append(value)
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        header = sheet[1]
        for cell in header:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10, color="222222")
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        widths = [14, 24, 36, 11, 38, 16, 16, 42, 24, 28, 28, 32, 22]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
        if rows:
            reference = f"A1:{sheet.cell(len(rows) + 1, len(HEADERS)).coordinate}"
            table = Table(displayName=f"{status.title()}MattersTable", ref=reference)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            sheet.add_table(table)
        sheet.auto_filter.ref = f"A1:{sheet.cell(max(1, len(rows) + 1), len(HEADERS)).coordinate}"
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.title = "Matter Tracker Projection"
    output.parent.mkdir(parents=True, exist_ok=True)
    staged = io.BytesIO()
    workbook.save(staged)
    create_private_file(output, staged.getvalue())
    return {"ok": True, "output": str(output), "read_only_source": True, "counts": counts,
            "sanitized": sanitized, "truncated": truncated,
            "omitted_deadline_fields": omitted_deadline_fields,
            "lossless_projection": not sanitized and not truncated and not omitted_deadline_fields}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate a disposable XLSX projection of the SQLite tracker")
    parser.add_argument("--database", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    try:
        emit(project(Path(args.database), Path(args.output)))
        return 0
    except Exception as exc:
        emit({"ok": False, "error": str(exc), "read_only_source": True})
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

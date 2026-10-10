from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from .common import canonical_json, sha256_bytes, sha256_file
from .document_helpers import extract_pdf


REVIEW_ACK_SCHEMA = "lawyeredos-content-review-ack-v1"
MAX_REVIEW_CHUNK_CHARS = 12_000
MAX_REVIEW_CONTENT_CHARS = 1_000_000
MAX_REVIEW_CHUNKS = 128
MAX_XLSX_CELLS = 250_000
DOCX_TEXT_PARTS = (
    "word/document.xml",
    "word/footnotes.xml",
    "word/endnotes.xml",
    "word/comments.xml",
)


def _unit(locator: dict[str, Any], text: str) -> dict[str, Any] | None:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    return None if not normalized else {"locator": locator, "text": normalized}


def _chunk_units(units: list[dict[str, Any]]) -> dict[str, Any]:
    source_characters = sum(len(item["text"]) for item in units)
    if source_characters > MAX_REVIEW_CONTENT_CHARS:
        return {
            "complete": False,
            "status": "extraction_limit",
            "characters": source_characters,
            "chunks": [],
            "reason": f"extracted content exceeds {MAX_REVIEW_CONTENT_CHARS} characters",
        }
    pieces: list[dict[str, Any]] = []
    for item in units:
        text = item["text"]
        for offset in range(0, len(text), MAX_REVIEW_CHUNK_CHARS):
            segment = text[offset:offset + MAX_REVIEW_CHUNK_CHARS]
            locator = {**item["locator"], "character_start": offset, "character_end": offset + len(segment)}
            pieces.append({"locator": locator, "text": segment})
    chunks: list[dict[str, Any]] = []
    current: list[dict[str, Any]] = []
    current_size = 0
    for piece in pieces:
        separator = 2 if current else 0
        if current and current_size + separator + len(piece["text"]) > MAX_REVIEW_CHUNK_CHARS:
            chunks.append(_render_chunk(len(chunks), current))
            current, current_size = [], 0
            separator = 0
        current.append(piece)
        current_size += separator + len(piece["text"])
    if current:
        chunks.append(_render_chunk(len(chunks), current))
    if len(chunks) > MAX_REVIEW_CHUNKS:
        return {
            "complete": False,
            "status": "chunk_limit",
            "characters": source_characters,
            "chunks": [],
            "reason": f"review requires more than {MAX_REVIEW_CHUNKS} bounded chunks",
        }
    return {
        "complete": True,
        "status": "extracted",
        "characters": source_characters,
        "chunks": chunks,
    }


def _render_chunk(index: int, pieces: list[dict[str, Any]]) -> dict[str, Any]:
    text = "\n\n".join(piece["text"] for piece in pieces)
    locators = [piece["locator"] for piece in pieces]
    return {
        "index": index,
        "text": text,
        "text_sha256": sha256_bytes(text.encode("utf-8")),
        "locators": locators,
    }


def message_review_content(text: str) -> dict[str, Any]:
    unit = _unit({"kind": "message-body"}, text)
    result = _chunk_units([] if unit is None else [unit])
    return {**result, "format": "message-body", "units": 0 if unit is None else 1}


def _docx_units(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    units: list[dict[str, Any]] = []
    reviewed_parts: list[str] = []
    visual_parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        parts = [name for name in DOCX_TEXT_PARTS if name in names]
        parts.extend(sorted(name for name in names if name.startswith("word/header") and name.endswith(".xml")))
        parts.extend(sorted(name for name in names if name.startswith("word/footer") and name.endswith(".xml")))
        for part_name in parts:
            root = ElementTree.fromstring(archive.read(part_name))
            paragraph_number = 0
            for paragraph in root.iter():
                if paragraph.tag.rsplit("}", 1)[-1] != "p":
                    continue
                paragraph_number += 1
                values: list[str] = []
                for element in paragraph.iter():
                    name = element.tag.rsplit("}", 1)[-1]
                    if name in {"t", "instrText"} and element.text:
                        values.append(element.text)
                    elif name == "delText" and element.text:
                        values.append(f"[deleted: {element.text}]")
                    elif name == "tab":
                        values.append("\t")
                    elif name in {"br", "cr"}:
                        values.append("\n")
                item = _unit(
                    {"kind": "docx", "part": part_name, "paragraph": paragraph_number},
                    "".join(values),
                )
                if item is not None:
                    units.append(item)
            reviewed_parts.append(part_name)
        visual_parts = sorted(
            name for name in names
            if name.startswith("word/media/") and not name.endswith("/")
        )
    return units, {
        "reviewed_parts": reviewed_parts,
        "visual_parts": visual_parts,
        "visual_review_required": bool(visual_parts),
    }


def _xlsx_units(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    units: list[dict[str, Any]] = []
    sheets: list[dict[str, Any]] = []
    examined_cells = 0
    try:
        for sheet in workbook.worksheets:
            dimension_cells = sheet.max_row * sheet.max_column
            if examined_cells + dimension_cells > MAX_XLSX_CELLS:
                return units, {
                    "sheets": sheets,
                    "cells_examined": examined_cells,
                    "complete": False,
                    "reason": f"workbook exceeds {MAX_XLSX_CELLS} bounded cells",
                }
            sheet_rows = 0
            nonempty_cells = 0
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                examined_cells += len(row)
                if examined_cells > MAX_XLSX_CELLS:
                    return units, {
                        "sheets": sheets,
                        "cells_examined": examined_cells,
                        "complete": False,
                        "reason": f"workbook exceeds {MAX_XLSX_CELLS} examined cells",
                    }
                values = []
                for cell in row:
                    if cell.value is not None:
                        values.append(f"{cell.coordinate}={cell.value}")
                    if cell.comment is not None:
                        values.append(
                            f"{cell.coordinate} comment by {cell.comment.author}: {cell.comment.text}"
                        )
                if not values:
                    continue
                sheet_rows += 1
                nonempty_cells += len(values)
                item = _unit(
                    {"kind": "xlsx", "sheet": sheet.title, "sheet_state": sheet.sheet_state, "row": row_number},
                    " | ".join(values),
                )
                if item is not None:
                    units.append(item)
            sheets.append({
                "title": sheet.title,
                "state": sheet.sheet_state,
                "rows_with_values": sheet_rows,
                "nonempty_cells": nonempty_cells,
                "visual_objects": len(getattr(sheet, "_images", ())) + len(getattr(sheet, "_charts", ())),
            })
    finally:
        workbook.close()
    visual_objects = sum(item["visual_objects"] for item in sheets)
    return units, {
        "sheets": sheets,
        "cells_examined": examined_cells,
        "visual_objects": visual_objects,
        "complete": visual_objects == 0,
        "reason": "workbook contains charts or images requiring separate visual review" if visual_objects else "",
    }


def _pdf_units(path: Path, scratch: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    scratch.mkdir(parents=True, exist_ok=True)
    output = scratch / f"{sha256_file(path)}.txt"
    if output.exists():
        text = output.read_text(encoding="utf-8", errors="replace")
        extraction = {"needs_ocr": len(text.strip()) < 40}
    else:
        extraction = extract_pdf(path, output)
        text = output.read_text(encoding="utf-8", errors="replace")
    pages = text.split("\f")
    while pages and not pages[-1].strip():
        pages.pop()
    units = [
        item for page_number, page in enumerate(pages, start=1)
        if (item := _unit({"kind": "pdf", "page": page_number}, page)) is not None
    ]
    empty_pages = [number for number, page in enumerate(pages, start=1) if not page.strip()]
    needs_ocr = bool(extraction.get("needs_ocr")) or bool(empty_pages) or not units
    return units, {
        "pages": len(pages),
        "empty_pages": empty_pages,
        "needs_ocr": needs_ocr,
        "complete": not needs_ocr,
        "reason": "one or more PDF pages require OCR or visual completeness review" if needs_ocr else "",
    }


def extract_review_content(path: Path, suffix: str, scratch: Path) -> dict[str, Any]:
    if suffix in {".txt", ".md", ".csv"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        units = [
            item for line_number, line in enumerate(text.splitlines(), start=1)
            if (item := _unit({"kind": suffix[1:], "line": line_number}, line)) is not None
        ]
        coverage: dict[str, Any] = {"lines": len(text.splitlines()), "complete": True}
    elif suffix == ".docx":
        units, coverage = _docx_units(path)
        coverage["complete"] = not coverage["visual_review_required"]
        if coverage["visual_review_required"]:
            coverage["reason"] = "DOCX contains embedded media requiring separate visual review"
    elif suffix == ".xlsx":
        units, coverage = _xlsx_units(path)
    elif suffix == ".pdf":
        units, coverage = _pdf_units(path, scratch)
    else:
        raise ValueError("unsupported review type")
    chunks = _chunk_units(units)
    complete = bool(coverage.get("complete", True) and chunks["complete"])
    reason = str(coverage.get("reason") or chunks.get("reason") or "")
    return {
        **chunks,
        "complete": complete,
        "status": "extracted" if complete else "extraction_incomplete",
        "format": suffix[1:],
        "units": len(units),
        "coverage": coverage,
        "reason": reason,
    }


def validate_acknowledgement(value: dict[str, Any], chunk_sha256: str) -> dict[str, Any]:
    if not isinstance(value, dict) or value.get("schema") != REVIEW_ACK_SCHEMA:
        raise ValueError(f"review acknowledgement must use schema {REVIEW_ACK_SCHEMA}")
    if value.get("reviewed_text_sha256") != chunk_sha256:
        raise ValueError("review acknowledgement hash does not match the pending chunk")
    reviewer = " ".join(str(value.get("reviewer", "")).split())
    summary = " ".join(str(value.get("summary", "")).split())
    evidence = value.get("evidence")
    limitations = value.get("limitations", [])
    if not reviewer or len(reviewer) > 100:
        raise ValueError("review acknowledgement requires a bounded reviewer identifier")
    if len(summary) < 10 or len(summary) > 2_000:
        raise ValueError("review acknowledgement requires a 10-2000 character summary")
    if not isinstance(evidence, list) or not 1 <= len(evidence) <= 20:
        raise ValueError("review acknowledgement requires 1-20 evidence observations")
    normalized_evidence: list[dict[str, str]] = []
    for item in evidence:
        if not isinstance(item, dict):
            raise ValueError("review evidence entries must be objects")
        locator = " ".join(str(item.get("locator", "")).split())
        observation = " ".join(str(item.get("observation", "")).split())
        if not locator or len(locator) > 500 or not observation or len(observation) > 1_000:
            raise ValueError("review evidence requires bounded locator and observation text")
        normalized_evidence.append({"locator": locator, "observation": observation})
    if not isinstance(limitations, list) or len(limitations) > 20:
        raise ValueError("review limitations must be a list of at most 20 strings")
    normalized_limitations = [" ".join(str(item).split()) for item in limitations]
    if any(not item or len(item) > 1_000 for item in normalized_limitations):
        raise ValueError("review limitations must contain bounded non-empty strings")
    if normalized_limitations:
        raise ValueError("a chunk with unresolved review limitations cannot be acknowledged complete")
    return {
        "schema": REVIEW_ACK_SCHEMA,
        "reviewed_text_sha256": chunk_sha256,
        "reviewer": reviewer,
        "summary": summary,
        "evidence": normalized_evidence,
        "limitations": [],
    }


def acknowledgement_digest(value: dict[str, Any]) -> str:
    return sha256_bytes(canonical_json(value).encode("utf-8"))

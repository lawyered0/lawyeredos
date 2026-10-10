from __future__ import annotations

import base64
import hashlib
import io
import json
import sqlite3
import tempfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from lawyeredos.mail_intake import SCHEMA, plan, process_bundle
from lawyeredos.tracker_cli import initialize


ROUTES = {
    "claude-gmail-connector": "claude-session-tool",
    "codex-session-tools": "codex-session-tool",
}


def _encoded(identifier: str, filename: str, mime_type: str, content: bytes) -> dict:
    return {
        "id": identifier,
        "filename": filename,
        "mime_type": mime_type,
        "size": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
        "content_base64": base64.b64encode(content).decode("ascii"),
    }


def _attachments() -> list[dict]:
    document_buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Northstar fictional DOCX evidence reviewed in full.")
    document.save(document_buffer)

    workbook_buffer = io.BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Fictional ledger"
    worksheet.append(["Item", "Amount"])
    worksheet.append(["Northstar flour", 125])
    workbook.save(workbook_buffer)

    return [
        _encoded("pdf-1", "fictional-evidence.pdf", "application/pdf", b"%PDF-1.4\n% fictional fixture\n"),
        _encoded(
            "docx-1", "fictional-letter.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            document_buffer.getvalue(),
        ),
        _encoded(
            "xlsx-1", "fictional-ledger.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            workbook_buffer.getvalue(),
        ),
        _encoded("txt-1", "fictional-note.txt", "text/plain", b"Northstar fictional TXT evidence.\n"),
        _encoded("md-1", "fictional-summary.md", "text/markdown", b"# Northstar\n\nFictional Markdown evidence.\n"),
        _encoded("csv-1", "fictional-data.csv", "text/csv", b"item,amount\nflour,125\n"),
    ]


def _message(message_id: str, attachments: list[dict] | None = None) -> dict:
    return {
        "id": message_id,
        "thread_id": f"thread-{message_id}",
        "received_at": "2032-02-01T12:00:00+00:00",
        "from": "client@example.invalid",
        "to": "lawyer@example.invalid",
        "cc": "",
        "subject": "Fictional supply dispute",
        "snippet": "Northstar Bakery asks counsel to review fictional evidence.",
        "body": (
            "Northstar Bakery Ltd. asks counsel to review this complete fictional message body "
            "about the fictional supply dispute. No real client, mailbox, or credential is used."
        ),
        "body_complete": True,
        "attachments": [] if attachments is None else attachments,
    }


def _bundle(planned: dict, source: str, fingerprint: str, messages: list[dict]) -> dict:
    return {
        "schema": SCHEMA,
        "source": source,
        "account_fingerprint": fingerprint,
        "query": planned["query"],
        "coverage": {"complete": True, "completed_at": "2032-02-01T13:00:00+00:00"},
        "messages": messages,
    }


def _checkpoint(database: Path) -> tuple[str, str, str, str, str] | None:
    with sqlite3.connect(database) as connection:
        return connection.execute(
            "SELECT file_number, account_fingerprint, completed_at, source, query FROM scan_state"
        ).fetchone()


def _fake_pdf_extract(_source: Path, output: Path) -> dict:
    output.write_text(
        "Northstar fictional PDF evidence was deterministically extracted for host-route acceptance.\n",
        encoding="utf-8",
    )
    return {"needs_ocr": False}


def _run_route(root: Path, adapter: str, source: str) -> dict:
    workspace = root / adapter
    matter = workspace / "2032-001-northstar"
    matter.mkdir(parents=True)
    (matter / "_matter-brief.md").write_text("# Fictional matter\n", encoding="utf-8")
    tracker = workspace / "matters.sqlite3"
    initialize(tracker)
    with sqlite3.connect(tracker) as connection:
        connection.execute(
            """INSERT INTO matters(file_number, client, description, matter_folder,
            status, opposing_party, primary_email, created_at, updated_at)
            VALUES (?, ?, ?, ?, 'open', ?, ?, ?, ?)""",
            (
                "2032-001", "Northstar Bakery Ltd.", "Fictional supply dispute",
                "2032-001-northstar", "Harbor Paper Co.",
                "client@example.invalid", "2032-01-01", "2032-01-01",
            ),
        )
    template = root / f"{adapter}-template.docx"
    template.write_bytes(b"fictional template placeholder")
    config = root / f"{adapter}.toml"
    config.write_text(
        f'''schema_version = 1
[workspace]
root = "{workspace}"
scratch = "{workspace / '.lawyeredos' / 'scratch'}"
timezone = "UTC"
[tracker]
database = "{tracker}"
[documents]
template = "{template}"
author = "Lawyered"
[jurisdiction]
profile = "none"
[integrations]
mail_adapter = "{adapter}"
calendar_adapter = "unavailable"
practice_manager_adapter = "unavailable"
[calendar]
calendar_id = ""
[review]
external_ai = false
''',
        encoding="utf-8",
    )
    fingerprint = hashlib.sha256(f"{source}:fictional@example.invalid".encode()).hexdigest()

    planned = plan(config, "2032-001")
    assert planned["source"] == source and planned["route"] == source
    all_types = _attachments()
    with patch("lawyeredos.mail_intake.extract_pdf", _fake_pdf_extract):
        first = process_bundle(
            config, "2032-001", _bundle(planned, source, fingerprint, [_message("all-types", all_types)])
        )
    assert first["status"] == "complete" and first["checkpoint_advanced"]
    reviewed = first["messages"][0]["attachments"]
    assert {Path(item["saved_path"]).suffix for item in reviewed} == {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv"}
    assert all(item["status"] == "reviewed" and item["review"]["text"] for item in reviewed)
    assert "Northstar fictional PDF evidence" in next(item for item in reviewed if item["original_name"].endswith(".pdf"))["review"]["text"]
    assert "DOCX evidence" in next(item for item in reviewed if item["original_name"].endswith(".docx"))["review"]["text"]
    assert "Northstar flour" in next(item for item in reviewed if item["original_name"].endswith(".xlsx"))["review"]["text"]

    repeated = process_bundle(
        config, "Northstar",
        _bundle(plan(config, "Northstar"), source, fingerprint, [_message("all-types", all_types)]),
    )
    assert repeated["messages"] == [{"message_id": "all-types", "status": "already_reviewed"}]
    assert len(list((matter / "email-attachments").iterdir())) == 6

    new_message = process_bundle(
        config, "2032-001",
        _bundle(plan(config, "2032-001"), source, fingerprint, [_message("new-message")]),
    )
    assert new_message["status"] == "complete" and new_message["messages"][0]["status"] == "reviewed"

    state = workspace / ".lawyeredos" / "data" / "mail-intake.sqlite3"
    checkpoint_before_partial = _checkpoint(state)
    missing = _encoded("retry-txt", "retry.txt", "text/plain", b"Complete fictional retry evidence.\n")
    missing.pop("content_base64")
    partial_message = _message("retry-message", [missing])
    partial_message["body_complete"] = False
    partial = process_bundle(
        config, "2032-001",
        _bundle(plan(config, "2032-001"), source, fingerprint, [partial_message]),
    )
    assert partial["status"] == "partial" and not partial["checkpoint_advanced"]
    assert partial["messages"][0]["attachments"][0]["status"] == "deferred"
    assert _checkpoint(state) == checkpoint_before_partial

    retry_attachment = _encoded("retry-txt", "retry.txt", "text/plain", b"Complete fictional retry evidence.\n")
    retry = process_bundle(
        config, "Northstar",
        _bundle(plan(config, "Northstar"), source, fingerprint, [_message("retry-message", [retry_attachment])]),
    )
    assert retry["status"] == "complete" and retry["checkpoint_advanced"]
    assert retry["messages"][0]["attachments"][0]["status"] == "reviewed"
    with sqlite3.connect(state) as connection:
        message_count = connection.execute("SELECT COUNT(*) FROM message_reviews").fetchone()[0]
        attachment_count = connection.execute("SELECT COUNT(*) FROM attachment_reviews").fetchone()[0]
    assert message_count == 3 and attachment_count == 7

    return {
        "adapter": adapter,
        "normalized_source": source,
        "full_body": "reviewed",
        "attachment_types": ["pdf", "docx", "xlsx", "txt", "md", "csv"],
        "repeat": "deduplicated",
        "new_message": "reviewed",
        "partial": "checkpoint_held",
        "retry": "completed",
        "message_reviews": message_count,
        "attachment_reviews": attachment_count,
    }


def run_acceptance() -> dict:
    with tempfile.TemporaryDirectory(prefix="lawyeredos-mail-host-acceptance-") as temporary:
        root = Path(temporary)
        routes = [_run_route(root, adapter, source) for adapter, source in ROUTES.items()]
    return {
        "ok": True,
        "fictional_only": True,
        "live_connector_used": False,
        "boundary": "normalized host bundle contract; no host tool name or live capability is assumed",
        "routes": routes,
    }


def main() -> int:
    print(json.dumps(run_acceptance(), indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

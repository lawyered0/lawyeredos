from __future__ import annotations

import argparse
import tempfile
import zipfile
from pathlib import Path

from scrub_docx_metadata import scrub


def require_docx():
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise SystemExit("python-docx is required; install requirements-documents.txt") from exc
    return Document, WD_ALIGN_PARAGRAPH, qn, Inches, Pt


def scrub_core_properties(document) -> None:
    props = document.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "Neutral correspondence template"
    props.subject = ""
    props.comments = ""
    props.keywords = ""
    props.category = ""


def build(output: Path) -> None:
    Document, WD_ALIGN_PARAGRAPH, qn, Inches, Pt = require_docx()
    if output.exists():
        raise ValueError("output already exists")
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("ORGANIZATION NAME")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    contact = header.add_run("\nStreet Address  |  City Region Postal Code  |  contact@example.invalid")
    contact.font.name = "Arial"
    contact.font.size = Pt(8)
    contact._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    contact._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    for text in ("{{DATE}}", "{{RECIPIENT}}", "{{ADDRESS}}"):
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(0)
    document.add_paragraph()
    subject = document.add_paragraph()
    subject.add_run("Re: ").bold = True
    subject.add_run("{{SUBJECT}}")
    subject.paragraph_format.space_after = Pt(12)
    body = document.add_paragraph("{{BODY}}")
    body.paragraph_format.space_after = Pt(10)
    document.add_paragraph("{{SIGNOFF}}")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("CONFIDENTIAL CORRESPONDENCE")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    footer_run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    scrub_core_properties(document)
    with tempfile.TemporaryDirectory(prefix=".neutral-template-", dir=output.parent) as temporary:
        staged = Path(temporary) / "template.docx"
        document.save(staged)
        scrub(staged, output)
    with zipfile.ZipFile(output) as package:
        names = set(package.namelist())
        if "docProps/custom.xml" in names:
            raise ValueError("unexpected custom properties in template")


def main() -> int:
    parser = argparse.ArgumentParser(description="Create a neutral correspondence template")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    try:
        build(Path(args.output).expanduser().resolve())
        return 0
    except (ValueError, OSError) as exc:
        print(f"error: {exc}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

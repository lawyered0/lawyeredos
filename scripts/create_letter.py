from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import stat
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree


PLACEHOLDER = re.compile(r"\{\{[A-Z][A-Z0-9_]*\}\}")
DRAFTING_MARKER = re.compile(
    r"\[(?:LAWYER\s+TO\s+CONFIRM|VERIFY|CITE|TBD|INSERT(?:[^\]]{0,80})?)\]"
    r"|(?<![A-Za-z0-9_])TK(?![A-Za-z0-9_])",
    re.IGNORECASE,
)
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REVISION_ELEMENTS = {
    "ins", "del", "moveFrom", "moveTo", "moveFromRangeStart",
    "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd", "cellIns",
    "cellDel", "cellMerge", "conflictIns", "conflictDel", "trackRevisions",
    "customXmlInsRangeStart", "customXmlInsRangeEnd",
    "customXmlDelRangeStart", "customXmlDelRangeEnd", "customXmlMoveFromRangeStart",
    "customXmlMoveFromRangeEnd", "customXmlMoveToRangeStart",
    "customXmlMoveToRangeEnd", "numberingChange", "pPrChange", "rPrChange",
    "sectPrChange", "tblPrChange", "tblPrExChange", "tblGridChange",
    "trPrChange", "tcPrChange",
}
COMMENT_ELEMENTS = {"commentRangeStart", "commentRangeEnd", "commentReference"}
MAX_DOCX_BYTES = 100 * 1024 * 1024
MAX_BODY_BYTES = 5 * 1024 * 1024
MAX_DOCX_ENTRY_BYTES = 50 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DOCX_ENTRIES = 10_000
REQUIRED_DOCX_PARTS = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
MAIN_DOCUMENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
OFFICE_DOCUMENT_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"


def require_docx():
    try:
        from docx import Document
        from docx.oxml import OxmlElement
    except ImportError as exc:
        raise SystemExit("python-docx is required; install requirements-documents.txt") from exc
    return Document, OxmlElement


def _paragraphs_in(parent):
    """Yield paragraphs in a document/header/footer, including table cells."""
    seen = set()

    def walk(container):
        for paragraph in container.paragraphs:
            marker = id(paragraph._p)
            if marker not in seen:
                seen.add(marker)
                yield paragraph
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(parent)


def _replace_in_runs(paragraph, key: str, value: str) -> None:
    """Replace a possibly split placeholder without flattening other runs."""
    search_from = 0
    while True:
        runs = list(paragraph.runs)
        combined = "".join(run.text for run in runs)
        start = combined.find(key, search_from)
        if start < 0:
            return
        end = start + len(key)
        offsets = []
        cursor = 0
        for run in runs:
            offsets.append((cursor, cursor + len(run.text)))
            cursor += len(run.text)
        touched = [i for i, (left, right) in enumerate(offsets) if left < end and right > start]
        if not touched:
            return
        first, last = touched[0], touched[-1]
        first_left, _ = offsets[first]
        last_left, _ = offsets[last]
        prefix = runs[first].text[:start - first_left]
        suffix = runs[last].text[end - last_left:]
        runs[first].text = prefix + value + (suffix if first == last else "")
        if first != last:
            for index in touched[1:-1]:
                runs[index].text = ""
            runs[last].text = suffix
        search_from = start + len(value)


def replace_simple(document, mapping: dict[str, str]) -> None:
    if any(PLACEHOLDER.search(value) for value in mapping.values()):
        raise ValueError("field values must not contain template placeholders")
    paragraphs = list(_paragraphs_in(document))
    for section in document.sections:
        paragraphs.extend(_paragraphs_in(section.header))
        paragraphs.extend(_paragraphs_in(section.footer))
    for paragraph in paragraphs:
        for key, value in mapping.items():
            _replace_in_runs(paragraph, key, value)


def replace_body(document, body_text: str, OxmlElement) -> None:
    target = next((paragraph for paragraph in document.paragraphs if "{{BODY}}" in paragraph.text), None)
    if target is None:
        raise ValueError("template is missing {{BODY}}")
    blocks = [block.strip() for block in re.split(r"\n\s*\n", body_text.strip()) if block.strip()]
    if not blocks:
        raise ValueError("body is empty")
    for block in blocks:
        paragraph_xml = OxmlElement("w:p")
        target._p.addprevious(paragraph_xml)
        from docx.text.paragraph import Paragraph
        paragraph = Paragraph(paragraph_xml, target._parent)
        paragraph.style = target.style
        paragraph.add_run(block)
    target._element.getparent().remove(target._element)


def scrub(document, author: str) -> None:
    props = document.core_properties
    props.author = author
    props.last_modified_by = author
    props.title = "Correspondence"
    props.subject = ""
    props.comments = ""
    props.keywords = ""
    props.category = ""


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _validate_docx_package(package: zipfile.ZipFile) -> None:
    infos = package.infolist()
    if len(infos) > MAX_DOCX_ENTRIES:
        raise ValueError("DOCX package contains too many entries")
    names = [info.filename for info in infos]
    if len(names) != len(set(names)):
        raise ValueError("DOCX package contains duplicate entries")
    if not REQUIRED_DOCX_PARTS.issubset(names):
        raise ValueError("input is not a valid Word DOCX package")
    total = 0
    for info in infos:
        parsed = PurePosixPath(info.filename)
        if parsed.is_absolute() or ".." in parsed.parts or "\\" in info.filename:
            raise ValueError("DOCX package contains an unsafe member name")
        if info.file_size > MAX_DOCX_ENTRY_BYTES:
            raise ValueError("DOCX package entry exceeds the uncompressed safety limit")
        total += info.file_size
        if total > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise ValueError("DOCX package exceeds the uncompressed safety limit")
    try:
        types = ElementTree.fromstring(package.read("[Content_Types].xml"))
        relationships = ElementTree.fromstring(package.read("_rels/.rels"))
        document = ElementTree.fromstring(package.read("word/document.xml"))
    except ElementTree.ParseError as exc:
        raise ValueError("input is not a valid Word DOCX package") from exc
    if types.tag != f"{{{CONTENT_TYPES_NS}}}Types":
        raise ValueError("input is not a valid Word DOCX package")
    overrides = {
        (item.get("PartName"), item.get("ContentType"))
        for item in types.findall(f"{{{CONTENT_TYPES_NS}}}Override")
    }
    if ("/word/document.xml", MAIN_DOCUMENT_TYPE) not in overrides:
        raise ValueError("input is not a valid Word DOCX package")
    if relationships.tag != f"{{{PACKAGE_REL_NS}}}Relationships":
        raise ValueError("input is not a valid Word DOCX package")
    office_targets = {
        item.get("Target", "").lstrip("/")
        for item in relationships.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
        if item.get("Type") == OFFICE_DOCUMENT_REL and not item.get("TargetMode")
    }
    if "word/document.xml" not in office_targets or document.tag != f"{{{WORD_NS}}}document":
        raise ValueError("input is not a valid Word DOCX package")


def package_review_markers(package: zipfile.ZipFile) -> list[str]:
    _validate_docx_package(package)
    markers = set()
    for name in package.namelist():
        if name.startswith("word/comments") and name.endswith(".xml"):
            markers.add("comments part")
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        root = ElementTree.fromstring(package.read(name))
        for element in root.iter():
            local = _local_name(element.tag)
            if element.tag.startswith(f"{{{WORD_NS}}}") and local in REVISION_ELEMENTS:
                markers.add(f"revision marker {local}")
            if element.tag.startswith(f"{{{WORD_NS}}}") and local in COMMENT_ELEMENTS:
                markers.add(f"comment marker {local}")
    return sorted(markers)


def package_placeholders(package: zipfile.ZipFile) -> list[str]:
    unresolved = set()
    for name in package.namelist():
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        root = ElementTree.fromstring(package.read(name))
        text = "".join(root.itertext())
        unresolved.update(PLACEHOLDER.findall(text))
        unresolved.update(match.group(0) for match in DRAFTING_MARKER.finditer(text))
    return sorted(unresolved)


def _reject_symlink_parents(path: Path) -> None:
    current = path.parent
    root = Path(path.anchor)
    while current != current.parent:
        if current.is_symlink():
            # macOS exposes root-owned aliases such as /var -> /private/var.
            # They are immutable to an ordinary caller; nested/user-owned links
            # remain forbidden.
            metadata = current.lstat()
            if current.parent != root or getattr(metadata, "st_uid", -1) != 0:
                raise ValueError(f"output parent must not contain a symlink: {current}")
        current = current.parent


def _read_regular_bytes(path: Path, label: str, maximum: int) -> bytes:
    flags = os.O_RDONLY
    if hasattr(os, "O_NOFOLLOW"):
        flags |= os.O_NOFOLLOW
    try:
        descriptor = os.open(path, flags)
    except OSError as exc:
        raise ValueError(f"{label} must be a regular, non-symlink file") from exc
    try:
        metadata = os.fstat(descriptor)
        if not stat.S_ISREG(metadata.st_mode):
            raise ValueError(f"{label} must be a regular, non-symlink file")
        if metadata.st_size > maximum:
            raise ValueError(f"{label} exceeds the {maximum}-byte safety limit")
        with os.fdopen(descriptor, "rb") as stream:
            descriptor = -1
            data = stream.read(maximum + 1)
            if len(data) > maximum:
                raise ValueError(f"{label} exceeds the {maximum}-byte safety limit")
            return data
    finally:
        if descriptor >= 0:
            os.close(descriptor)


def _validate_template_manifest(template: Path, template_bytes: bytes, manifest: Path) -> None:
    manifest_bytes = _read_regular_bytes(manifest, "template manifest", 64 * 1024)
    try:
        value = json.loads(manifest_bytes)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("template manifest is not valid JSON") from exc
    if not isinstance(value, dict) or value.get("schema_version") != 1:
        raise ValueError("unsupported template manifest schema")
    if value.get("algorithm") != "sha256":
        raise ValueError("unsupported template manifest algorithm")
    relative = value.get("template")
    expected = value.get("sha256")
    if (
        not isinstance(relative, str)
        or Path(relative).is_absolute()
        or ".." in Path(relative).parts
        or not isinstance(expected, str)
        or len(expected) != 64
        or any(character not in "0123456789abcdef" for character in expected)
    ):
        raise ValueError("template manifest contains invalid binding metadata")
    bound = (manifest.parent / relative).absolute()
    if template.absolute() != bound:
        raise ValueError("template path does not match the integrity manifest")
    if hashlib.sha256(template_bytes).hexdigest() != expected:
        raise ValueError("template integrity check failed")


def _unlink_created(path: Path, identity) -> None:
    try:
        current = path.stat(follow_symlinks=False)
        if (identity.st_dev, identity.st_ino) == (current.st_dev, current.st_ino):
            path.unlink()
    except OSError:
        pass


def _write_new(path: Path, data: bytes) -> None:
    _reject_symlink_parents(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    _reject_symlink_parents(path)
    descriptor, staged_name = tempfile.mkstemp(prefix=".lawyeredos-", dir=path.parent)
    staged = Path(staged_name)
    created = os.fstat(descriptor)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            descriptor = -1
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        try:
            os.link(staged, path, follow_symlinks=False)
        except FileExistsError as exc:
            raise ValueError("output already exists") from exc
    except BaseException:
        raise
    finally:
        if descriptor >= 0:
            os.close(descriptor)
        _unlink_created(staged, created)


def create(args: argparse.Namespace) -> None:
    Document, OxmlElement = require_docx()
    template_input = Path(args.template).expanduser()
    output = Path(args.output).expanduser().absolute()
    if output.exists() or output.is_symlink():
        raise ValueError("output already exists")
    template_bytes = _read_regular_bytes(template_input, "template", MAX_DOCX_BYTES)
    _validate_template_manifest(
        template_input,
        template_bytes,
        Path(args.template_manifest).expanduser(),
    )
    with zipfile.ZipFile(io.BytesIO(template_bytes)) as package:
        review_markers = package_review_markers(package)
    if review_markers:
        raise ValueError("template contains pre-existing revision/comment markup: " + ", ".join(review_markers))
    body = _read_regular_bytes(
        Path(args.body_file).expanduser(), "body file", MAX_BODY_BYTES,
    ).decode("utf-8")
    body_marker = PLACEHOLDER.search(body) or DRAFTING_MARKER.search(body)
    if body_marker:
        raise ValueError(f"body contains an unresolved drafting marker: {body_marker.group(0)}")
    document = Document(io.BytesIO(template_bytes))
    replace_body(document, body, OxmlElement)
    replace_simple(document, {
        "ORGANIZATION NAME": args.organization,
        "{{DATE}}": args.date,
        "{{RECIPIENT}}": args.recipient,
        "{{ADDRESS}}": args.address,
        "{{SUBJECT}}": args.subject,
        "{{SIGNOFF}}": args.signoff,
    })
    scrub(document, args.author)
    staged = io.BytesIO()
    document.save(staged)
    generated = staged.getvalue()
    with zipfile.ZipFile(io.BytesIO(generated)) as package:
        unresolved = package_placeholders(package)
        if unresolved:
            raise ValueError(f"unresolved template placeholders: {unresolved}")
        if "docProps/custom.xml" in package.namelist():
            raise ValueError("custom properties remain in output")
        if package_review_markers(package):
            raise ValueError("generated output contains revision/comment markup")
    _write_new(output, generated)


def main() -> int:
    parser = argparse.ArgumentParser(description="Create neutral DOCX correspondence")
    parser.add_argument("--template", required=True)
    parser.add_argument("--template-manifest", required=True)
    parser.add_argument("--body-file", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--organization", required=True)
    parser.add_argument("--date", required=True)
    parser.add_argument("--recipient", required=True)
    parser.add_argument("--address", required=True)
    parser.add_argument("--subject", required=True)
    parser.add_argument("--signoff", required=True)
    parser.add_argument("--author", default="Authorized reviewer")
    args = parser.parse_args()
    try:
        create(args)
        return 0
    except (ValueError, OSError, KeyError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        print(f"error: {exc}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
